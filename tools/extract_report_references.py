from docx import Document
from pypdf import PdfReader

pdf_path = "/Users/ashishtoppo/Downloads/NetScan Report.pdf"
docx_path = "/Users/ashishtoppo/Downloads/Network_Topology_Mapper_Internship_Report_Template.docx"

reader = PdfReader(pdf_path)
print("PDF pages:", len(reader.pages))
for index, page in enumerate(reader.pages[:10], 1):
    text = (page.extract_text() or "").replace("\n", " | ")
    print(f"PDF page {index}: {text[:1400]}")

document = Document(docx_path)
print("DOCX paragraphs:", len(document.paragraphs))
for index, paragraph in enumerate(document.paragraphs[:180], 1):
    text = paragraph.text.strip()
    if text:
        print(f"DOCX {index}: {text[:260]}")
