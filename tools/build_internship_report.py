from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/ashishtoppo/Documents/Nmap/Network_Topology_Mapper_Vulnerability_Analysis_Report.docx")
SCREENSHOTS = [
    Path("/var/folders/nd/13r_rlfj0sq8ybc8y66fwq5c0000gn/T/TemporaryItems/NSIRD_screencaptureui_XYvA3e/Screenshot 2026-07-08 at 4.11.31 PM.png"),
    Path("/var/folders/nd/13r_rlfj0sq8ybc8y66fwq5c0000gn/T/TemporaryItems/NSIRD_screencaptureui_NofdZ6/Screenshot 2026-07-09 at 6.45.28 PM.png"),
]

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 32, 43)
MUTED = RGBColor(92, 105, 112)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "EEF5FA"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=INK, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def set_table_borders(table, color="B8C6D1", size="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_FILL)
        set_cell_text(cell, header, bold=True, color=DARK_BLUE, size=10)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            set_cell_text(cells[idx], str(value), size=9.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "Network Topology Mapper & Vulnerability Analyzer"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], size=9, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        set_run_font(run, size=9, color=MUTED)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LOVELY PROFESSIONAL UNIVERSITY")
    set_run_font(r, size=15, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SUMMER TRAINING / INTERNSHIP REPORT")
    set_run_font(r, size=14, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("NETWORK TOPOLOGY MAPPER &\nVULNERABILITY ANALYZER")
    set_run_font(r, size=24, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("A SOC-style network discovery, topology mapping, and local vulnerability analysis system")
    set_run_font(r, size=12, color=MUTED, italic=True)

    table = doc.add_table(rows=6, cols=2)
    set_table_width(table, [2.0, 4.0])
    set_table_borders(table, color="D5DEE7")
    rows = [
        ("Submitted By", "______________________________"),
        ("Registration No.", "______________________________"),
        ("School", "School of Computer Science and Engineering"),
        ("Program / Section", "______________________________"),
        ("Faculty Guide", "______________________________"),
        ("Session", "2025 - 2026"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        shade_cell(row.cells[0], LIGHT_FILL)
        set_cell_text(row.cells[0], label, bold=True, color=DARK_BLUE)
        set_cell_text(row.cells[1], value)
    doc.add_page_break()


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=DARK_BLUE)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.4])
    set_table_borders(table, color="A9C7DA")
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT_FILL)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=10)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    doc.add_paragraph()


def add_section_page(doc, title, paragraphs, bullets=None, table=None):
    add_heading(doc, title, 1)
    for paragraph in paragraphs:
        add_para(doc, paragraph)
    if bullets:
        add_bullets(doc, bullets)
    if table:
        add_table(doc, table["headers"], table["rows"], table["widths"])
    doc.add_page_break()


def build_report():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_section_page(
        doc,
        "DECLARATION",
        [
            "I hereby declare that this internship/project report titled 'Network Topology Mapper & Vulnerability Analyzer' is my original work carried out during the summer training/internship period.",
            "The project has been prepared for academic learning and practical demonstration of cybersecurity concepts including network discovery, Nmap scanning, XML parsing, vulnerability analysis, and report generation.",
            "The work presented in this report has not been submitted previously for any degree or diploma. All references, tools, and technologies used in the project have been acknowledged in the bibliography section.",
            "\n\nStudent Signature: ____________________________\nDate: ____________________________",
        ],
    )

    add_section_page(
        doc,
        "CERTIFICATE",
        [
            "This is to certify that the project titled 'Network Topology Mapper & Vulnerability Analyzer' was successfully completed as part of the Summer Training/Internship requirements.",
            "The project involved the design and development of a local cybersecurity application capable of discovering hosts, identifying open ports and services, analyzing vulnerabilities, calculating risk scores, storing scan history, and displaying network topology in a professional dashboard.",
            "During the development of this project, the student demonstrated sincere effort, technical understanding, and practical interest in network security, software development, and vulnerability assessment.",
            "\n\nFaculty Guide: ____________________________        Student: ____________________________\nDate: ____________________________",
        ],
    )

    add_section_page(
        doc,
        "ACKNOWLEDGEMENT",
        [
            "I express my sincere gratitude to my faculty mentor, the School of Computer Science and Engineering, and Lovely Professional University for providing me with the opportunity to complete this internship project.",
            "I am thankful for the guidance, encouragement, and constructive suggestions received during the development of the project. The learning process helped me connect theoretical concepts from computer networks and cybersecurity with practical implementation.",
            "I also thank my friends and classmates for their support, feedback, and motivation during the project development and testing process.",
            "This project helped me improve my understanding of Nmap scanning, secure backend design, frontend dashboards, XML processing, risk analysis, and professional technical reporting.",
        ],
    )

    add_section_page(
        doc,
        "ABSTRACT",
        [
            "Network Topology Mapper & Vulnerability Analyzer is a local cybersecurity application designed to discover network devices, map topology, detect exposed services, analyze vulnerabilities, and generate risk-based reports.",
            "The application uses Nmap as the scanning engine and reads scan results in XML format. The backend parses this XML data and converts it into structured host, port, service, operating system, and script evidence. The analysis engine then checks for known risky configurations and local CVE patterns.",
            "The frontend presents the results in a SOC-style dashboard with cards, charts, an interactive topology graph, device inspector, live terminal, NSE script manager, scan history, filters, and export options.",
            "The main value of the project is that it converts complex Nmap output into a readable, visual, and actionable security assessment interface.",
        ],
    )

    toc_items = [
        "1. Introduction of Organization",
        "2. Summer Training / Internship Content Detail",
        "3. Project Overview",
        "4. Problem Statement",
        "5. Objectives",
        "6. Existing System",
        "7. Proposed System",
        "8. Technology Stack",
        "9. System Architecture",
        "10. Project Workflow",
        "11. Nmap Scan Profiles",
        "12. NSE Script Manager",
        "13. Local Vulnerability Analysis Engine",
        "14. CVE Matching Engine",
        "15. Risk Engine",
        "16. Dashboard Design",
        "17. Interactive Network Topology",
        "18. Device Inspector",
        "19. History, Search, Filters, and Exports",
        "20. Security and Ethical Considerations",
        "21. Implementation Details",
        "22. Testing and Validation",
        "23. Challenges and Solutions",
        "24. Limitations",
        "25. Future Scope",
        "26. System Snapshots",
        "27. Conclusion",
        "28. Bibliography",
    ]
    add_heading(doc, "TABLE OF CONTENTS", 1)
    for item in toc_items:
        add_para(doc, item)
    doc.add_page_break()

    add_section_page(
        doc,
        "1. INTRODUCTION OF ORGANIZATION",
        [
            "The internship focused on practical learning in the field of computer networks, cybersecurity, full-stack development, and secure software design. The learning environment emphasized the need to understand how real networks are discovered, monitored, and assessed for security risks.",
            "During the training period, the project work was aligned with real-world cybersecurity practices used by network administrators, Security Operations Center teams, and penetration testers.",
            "The project was developed to demonstrate how command-line security tools can be converted into a user-friendly application that supports visualization, reporting, and decision-making.",
        ],
        bullets=[
            "Domain: Cybersecurity and network assessment",
            "Role: Full-stack developer and security analysis learner",
            "Focus: Network discovery, vulnerability analysis, topology mapping, and reporting",
        ],
    )

    add_section_page(
        doc,
        "2. SUMMER TRAINING / INTERNSHIP CONTENT DETAIL",
        [
            "The training content covered both theoretical and practical topics. The theoretical part included computer networks, TCP/IP communication, ports, protocols, and vulnerability assessment concepts.",
            "The practical part involved building a working web application with a frontend, backend, local scanning service, XML parser, analysis engine, and reporting interface.",
            "The project also required understanding how Nmap works, how NSE scripts extend scans, and how scan evidence can be converted into useful security findings.",
        ],
        bullets=[
            "Computer network fundamentals and IP addressing",
            "Nmap scanning techniques and scan profiles",
            "NSE script usage and security checks",
            "XML parsing and structured data transformation",
            "React, TypeScript, Vite, Node.js, and Express.js",
            "Local rule-based vulnerability analysis and risk scoring",
        ],
    )

    add_section_page(
        doc,
        "3. PROJECT OVERVIEW",
        [
            "Network Topology Mapper & Vulnerability Analyzer is a local security assessment platform. It scans a target network, identifies live devices, detects open ports and services, and presents the results in a professional dashboard.",
            "The tool is designed for students, network administrators, SOC analysts, and penetration testing learners who need an easy way to understand network exposure.",
            "Instead of asking the user to manually read Nmap terminal output, the application organizes the results into visual sections such as metric cards, charts, topology graph, device inspector, vulnerability findings, and reports.",
        ],
    )

    add_section_page(
        doc,
        "4. PROBLEM STATEMENT",
        [
            "Many organizations and learners use command-line tools like Nmap to scan networks, but raw scan output can be difficult to read and interpret. Important findings may be missed when results are long, technical, or not visually organized.",
            "Security teams need a platform that can combine host discovery, port scanning, service detection, vulnerability analysis, risk scoring, and reporting in one place.",
            "The main problem addressed by this project is the lack of a simple, local, and visual tool that converts Nmap scan results into clear security intelligence without depending on cloud AI services.",
        ],
        bullets=[
            "Unknown devices may exist in a network.",
            "Open ports and services may expose attack surfaces.",
            "Manual scan interpretation is slow and error-prone.",
            "Security reports are often difficult to generate from raw command output.",
        ],
    )

    add_section_page(
        doc,
        "5. OBJECTIVES",
        [
            "The objective of the project is to develop a complete local cybersecurity application that improves network visibility and vulnerability understanding.",
            "The project focuses on making Nmap-based scanning more accessible through a graphical dashboard and structured analysis pipeline.",
        ],
        bullets=[
            "Discover hosts, IP addresses, MAC vendors, and operating system details.",
            "Detect open ports, services, service versions, and NSE script evidence.",
            "Visualize devices in an interactive network topology map.",
            "Analyze risky services using local rules and CVE matching.",
            "Calculate device-level and network-level risk scores.",
            "Store scan history and support comparison between scans.",
            "Export reports in HTML, JSON, CSV, XML, PNG, and SVG formats.",
        ],
    )

    add_section_page(
        doc,
        "6. EXISTING SYSTEM",
        [
            "Traditional network scanning is usually performed through command-line tools. These tools are powerful but require technical knowledge and manual interpretation.",
            "A beginner may find it difficult to understand the difference between scan types, service detection, version detection, NSE scripts, and vulnerability evidence.",
            "Some web-based tools depend on cloud services, which may not be suitable for private LAN scanning because internal devices are not reachable from outside the network.",
            "The existing approach also lacks a unified dashboard where topology, vulnerabilities, risks, and reports are shown together.",
        ],
    )

    add_section_page(
        doc,
        "7. PROPOSED SYSTEM",
        [
            "The proposed system is a local full-stack application that runs Nmap on the user's machine and processes results locally. This helps scan private networks without sending scan data to external AI APIs or cloud scanners.",
            "The application provides predefined scan profiles, selected NSE scripts, live scan terminal output, XML parsing, local rule-based vulnerability detection, CVE matching, risk scoring, dashboards, topology visualization, history, and exports.",
            "By combining scanning, analysis, and visualization, the proposed system improves usability and reduces the time required to understand network security posture.",
        ],
        table={
            "headers": ["Area", "Proposed Improvement"],
            "rows": [
                ("Scanning", "Nmap profiles and NSE script selection from the UI"),
                ("Analysis", "Local rule engine, CVE matching, and risk scoring"),
                ("Visualization", "Dashboard, charts, topology graph, and device inspector"),
                ("Reporting", "HTML, JSON, CSV, XML, PNG, and SVG export support"),
                ("Privacy", "Offline analysis with no AI API dependency"),
            ],
            "widths": [1.6, 4.7],
        },
    )

    add_section_page(
        doc,
        "8. TECHNOLOGY STACK",
        [
            "The project uses a modern full-stack TypeScript architecture. Each technology was selected based on its role in the scanning, processing, visualization, and reporting workflow.",
        ],
        table={
            "headers": ["Technology", "Purpose"],
            "rows": [
                ("React", "Builds the interactive SOC-style user interface"),
                ("TypeScript", "Adds strict typing and reduces development errors"),
                ("Vite", "Provides fast development server and optimized frontend build"),
                ("Node.js", "Runs the backend service locally"),
                ("Express.js", "Creates API endpoints for scans, history, scripts, and exports"),
                ("Nmap", "Performs host discovery, port scanning, and service detection"),
                ("NSE Scripts", "Adds deeper checks for services and vulnerabilities"),
                ("XML Parsing", "Converts Nmap XML into structured application data"),
                ("React Flow", "Displays interactive network topology"),
                ("Recharts", "Displays security metrics and charts"),
                ("Local JSON History", "Stores scan records without requiring a database server"),
            ],
            "widths": [1.8, 4.5],
        },
    )

    add_section_page(
        doc,
        "9. SYSTEM ARCHITECTURE",
        [
            "The system follows a clear client-server architecture. The frontend is responsible for user interaction and visualization, while the backend performs scanning, parsing, analysis, and history management.",
            "The backend executes Nmap using safe argument arrays rather than shell strings. This reduces command-injection risk and keeps scan execution controlled.",
            "The architecture separates scanning, XML parsing, vulnerability analysis, risk scoring, report generation, and UI rendering into modular components.",
        ],
    )
    add_heading(doc, "Architecture Flow", 2)
    add_table(
        doc,
        ["Step", "Component", "Role"],
        [
            ("1", "React UI", "User enters target, profile, and scripts"),
            ("2", "Express API", "Receives scan request and validates input"),
            ("3", "Nmap Service", "Runs Nmap locally and collects XML output"),
            ("4", "XML Parser", "Extracts hosts, ports, services, OS, and script evidence"),
            ("5", "Analysis Engine", "Applies rules, CVE matching, and risk scoring"),
            ("6", "History Service", "Stores scan records locally"),
            ("7", "Dashboard", "Displays topology, charts, findings, and reports"),
        ],
        [0.7, 1.6, 4.0],
    )
    doc.add_page_break()

    add_section_page(
        doc,
        "10. PROJECT WORKFLOW",
        [
            "The workflow starts when the user enters a target such as a single IP address, hostname, or CIDR range. The user then selects a scan profile and optional NSE scripts.",
            "The frontend sends the request to the backend. The backend validates the input, builds the Nmap command, and starts the scan. While the scan is running, live terminal events are streamed back to the frontend.",
            "After completion, Nmap XML output is parsed and analyzed. The final scan record is saved and displayed in the dashboard.",
        ],
        bullets=[
            "Input target and select scan profile.",
            "Run Nmap locally through the backend service.",
            "Stream scan progress to the live terminal.",
            "Parse XML output into structured host data.",
            "Run local vulnerability and CVE analysis.",
            "Generate risk metrics and topology graph.",
            "Store scan history and export reports.",
        ],
    )

    add_section_page(
        doc,
        "11. NMAP SCAN PROFILES",
        [
            "The application provides predefined Nmap scan profiles so users do not need to remember command-line options. Each profile maps to a specific Nmap argument set.",
        ],
        table={
            "headers": ["Profile", "Nmap Options", "Purpose"],
            "rows": [
                ("Quick", "-T4 -F", "Fast scan of common ports"),
                ("Intense", "-T4 -A -v", "Detailed scan with OS, services, scripts, and verbosity"),
                ("OS Detection", "-O", "Attempts to identify operating systems"),
                ("Service Detection", "-sV", "Identifies services running on ports"),
                ("Version Detection", "-sV --version-all", "Performs deeper version detection"),
                ("UDP", "-sU --top-ports 200", "Scans common UDP services"),
                ("TCP SYN", "-sS", "Performs half-open TCP SYN scan"),
                ("ACK", "-sA", "Tests firewall filtering behavior"),
                ("NULL / FIN / XMAS", "-sN / -sF / -sX", "Firewall and TCP behavior testing"),
                ("Stealth", "-sS -T2 --data-length 24", "Slower scan with additional padding"),
                ("Aggressive", "-A -T4", "Comprehensive scan for rich results"),
                ("Ping Sweep", "-sn", "Discovers live hosts without port scanning"),
                ("Traceroute", "--traceroute", "Shows network path to target"),
            ],
            "widths": [1.4, 1.8, 3.1],
        },
    )

    add_section_page(
        doc,
        "12. NSE SCRIPT MANAGER",
        [
            "NSE stands for Nmap Scripting Engine. NSE scripts are small programs that extend Nmap beyond simple port scanning.",
            "The project includes an NSE Script Manager that loads local Nmap scripts, groups them into categories, and allows the user to enable or disable scripts before scanning.",
            "Selected scripts are appended to the scan command using the Nmap --script option. The script output is later parsed as evidence for vulnerability analysis.",
        ],
        table={
            "headers": ["Script", "Purpose"],
            "rows": [
                ("http-title", "Identifies web page title for quick service recognition"),
                ("http-security-headers", "Checks missing or weak HTTP security headers"),
                ("ssl-enum-ciphers", "Audits SSL/TLS versions and cipher strength"),
                ("ssh2-enum-algos", "Lists SSH algorithms and weak cryptographic choices"),
                ("ftp-anon", "Checks whether anonymous FTP login is allowed"),
                ("smb-protocols", "Detects SMB protocol versions including SMBv1"),
                ("smb-vuln-ms17-010", "Checks for MS17-010 SMB exposure"),
                ("snmp-info", "Collects SNMP device information when exposed"),
                ("dns-recursion", "Checks open DNS recursion"),
                ("mongodb-info", "Collects MongoDB service information"),
                ("redis-info", "Collects Redis service and authentication information"),
            ],
            "widths": [2.2, 4.1],
        },
    )

    add_section_page(
        doc,
        "13. LOCAL VULNERABILITY ANALYSIS ENGINE",
        [
            "The vulnerability analysis engine is local and rule-based. It does not call any external AI service. It checks scan evidence against predefined security rules.",
            "For example, if FTP anonymous login is detected, the engine marks the finding as high severity and provides a reason, impact, mitigation, and references.",
            "The rule engine helps convert technical scan evidence into actionable security findings that can be understood by students, administrators, and security teams.",
        ],
        bullets=[
            "Telnet exposure is marked risky because it sends data in plain text.",
            "SMBv1 exposure is treated as critical due to historical remote code execution issues.",
            "Open Redis and MongoDB are highlighted because they may expose sensitive data.",
            "Weak SSH or old TLS protocols are reported as cryptographic weaknesses.",
            "SNMP exposure is reported because it can reveal network inventory details.",
        ],
    )

    add_section_page(
        doc,
        "14. CVE MATCHING ENGINE",
        [
            "The CVE engine compares detected services, product names, versions, and operating system details with a local CVE-style database.",
            "A CVE is a public identifier for a known cybersecurity vulnerability. By matching service versions against known patterns, the application can identify possible vulnerable components.",
            "The CVE output includes CVE ID, CVSS score, severity, description, references, patch availability, and exploit availability where available.",
        ],
        table={
            "headers": ["Field", "Meaning"],
            "rows": [
                ("CVE ID", "Unique identifier for a known vulnerability"),
                ("CVSS Score", "Numerical severity score from 0 to 10"),
                ("Severity", "Critical, High, Medium, Low, or Informational"),
                ("Description", "Short explanation of the vulnerability"),
                ("References", "Links to trusted vulnerability details"),
                ("Patch Availability", "Indicates whether a fix is known"),
                ("Exploit Availability", "Indicates whether exploitation is publicly known"),
            ],
            "widths": [1.8, 4.5],
        },
    )

    add_section_page(
        doc,
        "15. RISK ENGINE",
        [
            "The risk engine calculates risk scores at both device level and network level. The score is based on open ports, exposed management services, number of vulnerabilities, and severity of findings.",
            "The purpose of risk scoring is to help users prioritize. A network may have many findings, but the highest-risk devices should be reviewed first.",
            "The dashboard shows metrics such as overall network score, security score, average risk, attack surface score, critical devices, severity distribution, and top exposed services.",
        ],
        bullets=[
            "High-risk devices are highlighted in the topology graph.",
            "Critical vulnerabilities increase the risk score more than low-severity findings.",
            "Remote access services such as RDP, SMB, Telnet, SSH, Redis, and MongoDB increase attack surface.",
            "Security score helps summarize the overall condition of the network.",
        ],
    )

    add_section_page(
        doc,
        "16. DASHBOARD DESIGN",
        [
            "The dashboard is designed with a professional SOC-style interface. It uses dark panels, clear metric cards, charts, terminal output, and structured controls.",
            "The first screen gives a quick security summary: total hosts, active hosts, critical hosts, open ports, detected vulnerabilities, services, average risk, and security score.",
            "Charts help users understand severity distribution, top ports, services, operating systems, and timeline activity.",
        ],
        bullets=[
            "Metric cards provide instant summary.",
            "Charts convert scan data into visual patterns.",
            "Live terminal improves transparency during scans.",
            "Filters and search help investigate large results.",
        ],
    )

    add_section_page(
        doc,
        "17. INTERACTIVE NETWORK TOPOLOGY",
        [
            "The topology graph is one of the main features of the project. It displays discovered devices as nodes and connections as edges.",
            "The graph supports zooming, panning, dragging, minimap navigation, and risk highlighting. Devices with higher risk scores are visually emphasized.",
            "The topology view helps users understand the structure of the network instead of only reading tabular scan results.",
        ],
        bullets=[
            "Gateway or firewall devices can act as central nodes.",
            "Hosts are grouped visually around the network structure.",
            "Critical or vulnerable hosts are easier to locate.",
            "The graph supports interactive investigation.",
        ],
    )

    add_section_page(
        doc,
        "18. DEVICE INSPECTOR",
        [
            "When a user clicks a device in the topology graph or device table, the Device Inspector panel opens.",
            "The inspector provides detailed information about the selected host, including IP address, hostname, MAC address, vendor, operating system, response time, TTL, ports, services, vulnerabilities, risk score, attack surface, and timeline.",
            "This panel helps users move from a high-level network view to a focused device-level investigation.",
        ],
        bullets=[
            "Shows host identity and classification.",
            "Lists open ports and detected services.",
            "Displays vulnerability findings with mitigation advice.",
            "Shows risk score and attack surface score.",
        ],
    )

    add_section_page(
        doc,
        "19. HISTORY, SEARCH, FILTERS, AND EXPORTS",
        [
            "The project stores each scan in local history. This allows users to review old scans and compare them with new scans.",
            "Search and filters make it easier to find hosts by IP, hostname, MAC address, vendor, operating system, service, port, severity, and risk score.",
            "Export features support professional reporting and sharing of results.",
        ],
        table={
            "headers": ["Feature", "Description"],
            "rows": [
                ("History", "Stores scan results for later review"),
                ("Compare Scans", "Identifies new hosts, removed hosts, new vulnerabilities, and resolved vulnerabilities"),
                ("Search", "Finds devices, services, ports, vendors, and vulnerabilities"),
                ("Filters", "Narrows results by subnet, OS, vendor, service, severity, and risk"),
                ("Exports", "Supports HTML, JSON, CSV, XML, PNG, and SVG outputs"),
            ],
            "widths": [1.7, 4.6],
        },
    )

    add_section_page(
        doc,
        "20. SECURITY AND ETHICAL CONSIDERATIONS",
        [
            "Network scanning must be performed responsibly. The tool is intended for authorized networks, academic labs, personal systems, and approved security testing environments.",
            "Unauthorized scanning can violate organizational policy or legal rules. Therefore, users should only scan devices and networks where they have explicit permission.",
            "The application avoids cloud AI APIs and processes scan data locally, which helps protect sensitive network information.",
        ],
        bullets=[
            "Scan only authorized systems.",
            "Avoid intrusive or brute-force NSE scripts unless permission is granted.",
            "Use aggressive scans carefully because they can create network noise.",
            "Protect exported reports because they contain sensitive security information.",
        ],
    )

    add_section_page(
        doc,
        "21. IMPLEMENTATION DETAILS",
        [
            "The project is implemented as a full-stack TypeScript application. The frontend uses React components, reusable hooks, context for settings, and service functions for API communication.",
            "The backend is organized into services such as scan manager, Nmap profiles, NSE service, XML parser, analysis engine, history service, report service, and system status service.",
            "Nmap is executed through the backend using argument arrays. Scan events are streamed to the frontend using server-sent events, allowing the live terminal to update during scan execution.",
        ],
    )
    add_table(
        doc,
        ["Module", "Responsibility"],
        [
            ("Scan Manager", "Starts scans and streams events"),
            ("Nmap Profiles", "Maps UI scan buttons to Nmap arguments"),
            ("NSE Service", "Loads and categorizes local NSE scripts"),
            ("XML Parser", "Parses Nmap XML into device objects"),
            ("Analysis Engine", "Runs rules, CVE matching, and risk scoring"),
            ("History Service", "Saves and loads scan records"),
            ("Report Service", "Exports reports in multiple formats"),
            ("React UI", "Displays dashboard, topology, terminal, and controls"),
        ],
        [1.8, 4.5],
    )
    doc.add_page_break()

    add_section_page(
        doc,
        "22. TESTING AND VALIDATION",
        [
            "Testing was performed by running the build process, checking API endpoints, validating Nmap availability, and verifying that the dashboard renders correctly.",
            "The application was tested with sample Nmap XML data and local Nmap installation. The sample data includes hosts, firewall, server, web application, IoT device, Redis, MongoDB, SMB, SSL, SSH, SNMP, and web security findings.",
            "The XML import feature was also validated to ensure saved Nmap XML files can be analyzed without running a new scan.",
        ],
        bullets=[
            "Build validation using TypeScript and Vite production build.",
            "API validation for system status, scan history, and NSE scripts.",
            "Sample scan validation using bundled Nmap XML fixture.",
            "Browser smoke testing for dashboard, topology, terminal, and charts.",
            "Manual review of scan control behavior and script selection.",
        ],
    )

    add_section_page(
        doc,
        "23. CHALLENGES AND SOLUTIONS",
        [
            "One challenge was converting Nmap output into useful structured data. This was solved by using XML output and a parser that extracts hosts, ports, services, operating systems, MAC vendors, and script evidence.",
            "Another challenge was real-time scan feedback. Waiting silently for a scan to finish would reduce user confidence. This was solved by streaming scan events to the live terminal.",
            "A further challenge was handling many local NSE scripts. A real Nmap installation can include hundreds of scripts, so the application uses a curated default selection to avoid overly heavy scans.",
            "The project also required balancing a visually impressive dashboard with readability and performance. The UI was organized into reusable panels, tables, filters, and graph components.",
        ],
        table={
            "headers": ["Challenge", "Solution"],
            "rows": [
                ("Raw scan output is complex", "Use XML parsing and structured device models"),
                ("Users need live feedback", "Stream scan events to a terminal panel"),
                ("Too many NSE scripts", "Use curated defaults and searchable script manager"),
                ("Need offline analysis", "Build local rule, CVE, and risk engines"),
                ("Need visual clarity", "Use dashboard cards, charts, topology, and inspector"),
            ],
            "widths": [2.1, 4.2],
        },
    )

    add_section_page(
        doc,
        "24. LIMITATIONS",
        [
            "The project depends on local Nmap installation for real scanning. If Nmap is not installed, the application can still import XML and replay sample data, but real scans require Nmap on the system path.",
            "Some Nmap scan types require administrator or root privileges. For example, SYN scan, OS detection, and some UDP scans may not work fully without elevated permissions.",
            "The CVE database included in the project is local and curated for demonstration. A production system would require regular updates from trusted vulnerability sources.",
            "Large enterprise scans may require job queues, database storage, pagination, and stronger resource management.",
        ],
    )

    add_section_page(
        doc,
        "25. FUTURE SCOPE",
        [
            "The project can be extended with several advanced cybersecurity and enterprise features.",
            "Future versions can include scheduled scans, email alerts, authentication, role-based access control, larger CVE feeds, CISA KEV mapping, MITRE ATT&CK mapping, compliance summaries, and passive discovery.",
            "For scalability, scan jobs can be moved into a background queue and stored in a database such as PostgreSQL. Large tables can use virtualization, and the frontend bundle can be split into smaller chunks.",
        ],
        bullets=[
            "Scheduled scans and automatic recurring monitoring.",
            "Role-based login system for teams.",
            "Database-backed history for large deployments.",
            "Updated CVE and CISA KEV feeds.",
            "MITRE ATT&CK and OWASP mapping.",
            "PDF report designer and branded templates.",
            "Real-time alerts and notification channels.",
            "Passive network discovery and inventory tracking.",
        ],
    )

    add_heading(doc, "26. SYSTEM SNAPSHOTS", 1)
    add_para(doc, "This section includes representative screenshots and descriptions of the developed system interface.")
    for idx, path in enumerate(SCREENSHOTS, 1):
        if path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(path), width=Inches(5.8))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = caption.add_run(f"Figure {idx}: Application screen showing scan controls and SOC dashboard elements")
            set_run_font(r, size=9.5, color=MUTED, italic=True)
            doc.add_paragraph()
    add_callout(
        doc,
        "Snapshot Note",
        "The interface is designed to look like a professional security operations console, with scan controls, status badges, live terminal output, topology views, and investigation panels."
    )
    doc.add_page_break()

    add_section_page(
        doc,
        "27. CONCLUSION",
        [
            "Network Topology Mapper & Vulnerability Analyzer successfully demonstrates how Nmap-based scanning can be converted into a complete cybersecurity dashboard.",
            "The project discovers network devices, detects ports and services, analyzes vulnerabilities, calculates risk scores, visualizes topology, stores scan history, and exports reports.",
            "The main achievement is that the application makes complex scan data easier to understand and act upon. It is useful for academic learning, network administration, and practical cybersecurity demonstrations.",
            "Overall, the project helped strengthen understanding of networking, secure backend development, frontend interface design, XML parsing, local analysis engines, and professional reporting.",
        ],
    )

    add_heading(doc, "28. BIBLIOGRAPHY", 1)
    bibliography = [
        "Nmap Official Website: https://nmap.org/",
        "Nmap Scripting Engine Documentation: https://nmap.org/book/nse.html",
        "National Vulnerability Database: https://nvd.nist.gov/",
        "OWASP Foundation: https://owasp.org/",
        "MITRE ATT&CK Framework: https://attack.mitre.org/",
        "React Documentation: https://react.dev/",
        "TypeScript Documentation: https://www.typescriptlang.org/",
        "Vite Documentation: https://vite.dev/",
        "Express.js Documentation: https://expressjs.com/",
        "React Flow Documentation: https://reactflow.dev/",
        "Recharts Documentation: https://recharts.org/",
    ]
    add_bullets(doc, bibliography)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_report()
    print(OUT)
