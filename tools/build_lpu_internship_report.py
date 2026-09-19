from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("output/documents/Network_Topology_Mapper_Vulnerability_Analyzer_LPU_Internship_Report.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(80, 80, 80)
LIGHT_FILL = "F2F4F7"
CALIBRI = "Calibri"


def set_run_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None,
                 color: RGBColor | None = None, name: str = CALIBRI) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            cell = row.cells[index]
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = CALIBRI
    normal._element.rPr.rFonts.set(qn("w:ascii"), CALIBRI)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), CALIBRI)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = CALIBRI
        style._element.rPr.rFonts.set(qn("w:ascii"), CALIBRI)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), CALIBRI)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Number", "List Bullet"]:
        style = styles[name]
        style.font.name = CALIBRI
        style._element.rPr.rFonts.set(qn("w:ascii"), CALIBRI)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), CALIBRI)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_footer(section, label: str) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(f"{label} | Page ")
    set_run_font(run, size=9, color=MUTED)
    add_page_number(paragraph)


def add_para(doc: Document, text: str, *, style: str | None = None, align=None, bold=False,
             italic=False, size: float | None = None, color: RGBColor | None = None,
             before: float | None = None, after: float | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    if before is not None:
        paragraph.paragraph_format.space_before = Pt(before)
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)


def add_body(doc: Document, text: str) -> None:
    for para in [p.strip() for p in text.strip().split("\n\n") if p.strip()]:
        add_para(doc, para)


def add_ascii_block(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text.strip("\n"))
    set_run_font(run, name="Courier New", size=9.5, color=BLACK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_FILL)
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, bold=True, size=10.5, color=BLACK)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=10)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    set_table_geometry(table, widths)
    doc.add_paragraph()


def words(text: str) -> int:
    return len(re.findall(r"\b[\w']+\b", text))


def paragraph_from_points(topic: str, points: list[str], closing: str) -> str:
    body = " ".join(points)
    return f"{topic} {body} {closing}"


ORG_SECTION = """
The internship was undertaken in the context of CipherShield Cybersecurity Solutions Pvt. Ltd., a realistic mid-sized cybersecurity services organization that focuses on practical defensive security for educational institutions, small enterprises, and technology-driven businesses. The company operates as a security engineering and consulting unit where network assessment, application security, log review, secure configuration, and incident readiness are treated as connected parts of one defensive program. Its work is not limited to running automated scanners; it includes understanding client environments, identifying exposed assets, documenting risk in language that administrators can act upon, and recommending controls that are realistic for the client's budget and technical maturity. This environment was suitable for the Network Topology Mapper & Vulnerability Analyzer because the project required both software development skill and security assessment judgment.

The mission of the organization is to help clients build measurable cyber resilience. In practical terms, this means giving organizations visibility into the systems they own, the services they expose, the weaknesses that can be exploited, and the sequence in which remediation should be performed. Many organizations have purchased security tools but still struggle to convert raw output into decisions. CipherShield therefore emphasizes clarity, evidence, prioritization, and repeatable processes. Its analysts are expected to explain why a finding matters, how it was detected, which asset is affected, and what technical action can reduce the risk. This mission directly shaped the project, because the application was designed not merely to run Nmap but to convert Nmap evidence into topology, severity, attack surface, and reportable security information.

The vision of CipherShield is to make defensive security understandable and continuous rather than occasional and reactive. The organization promotes the idea that security teams should know their network before an attacker studies it. A topology map, an inventory of reachable services, a list of known CVE matches, and a history of previous scans are foundational elements of this vision. The project supported this vision by combining discovery, vulnerability analysis, and visualization in a single browser-based console. It demonstrated how a student project can reflect real security operations center practices: scan, normalize evidence, classify risk, present a dashboard, and preserve results for later comparison.

The company provides services across several cybersecurity domains. Its network security practice includes host discovery, port scanning, firewall review, segmentation assessment, and vulnerability validation. Its application security practice focuses on input validation, authentication design, OWASP awareness, API testing, and secure coding review. Its governance and compliance support includes technical documentation, audit preparation, asset inventory, and remediation tracking. Its incident readiness work includes tabletop exercises, log source review, and basic containment planning. These service areas gave the internship a broad learning base. The project drew most heavily from network security, vulnerability management, secure coding, and reporting, but it also touched governance because a security tool is useful only when its output can be documented and defended.

During the internship, the assigned role was similar to a junior cybersecurity engineering intern. The responsibilities included studying Nmap behavior, understanding how XML output represents hosts and ports, designing a backend service that can execute scans safely, building a frontend dashboard, and preparing security reports that can be understood by both technical and academic reviewers. The role also required thinking about limitations. Network scanning can be intrusive, may require administrative privileges, and can produce false positives if evidence is interpreted without context. The internship therefore emphasized disciplined implementation: validate input, restrict dangerous assumptions, expose scan progress clearly, and label results as assessment evidence rather than absolute truth.

The learning experience was valuable because it connected theory with implementation. Networking concepts such as IP addressing, ports, protocols, TCP handshakes, and service banners became more meaningful when they were extracted from scan output and displayed as nodes, edges, and risk indicators. Web development concepts such as React components, TypeScript interfaces, Express routes, and JSON-based storage became more meaningful when they supported a cybersecurity workflow. The internship also developed report writing discipline. A vulnerability report should not be a collection of alarming statements; it should explain affected assets, exposure, evidence, business impact, severity, and mitigation. Through this project, the organization setting provided a realistic frame in which software engineering, cybersecurity analysis, and professional documentation could be practiced together.
"""


INTERNSHIP_CONTENT = """
Network fundamentals formed the foundation of the internship. A network is a collection of devices that communicate using protocols, addresses, and shared rules for transmitting data. From a security perspective, every reachable device, open port, exposed management interface, and unpatched service contributes to the observable attack surface. Understanding this foundation was necessary before implementing any scanner or topology tool. The project required the system to recognize hosts, group devices into subnets, infer device types, and represent communication relationships in a form that a user could understand. Network fundamentals also clarified why inventory is the first step in security assessment. An organization cannot protect systems that it has not identified, and it cannot prioritize remediation without knowing which services are exposed.

The TCP/IP model was studied as the practical protocol stack used on modern networks. The link layer handles local network communication, the internet layer provides addressing and routing through IP, the transport layer uses protocols such as TCP and UDP, and the application layer hosts services such as HTTP, SSH, DNS, SMB, FTP, and SNMP. This model helped in interpreting Nmap results because Nmap reports transport protocols, port states, service names, versions, and sometimes application fingerprints. TCP behavior was especially important because many scan techniques depend on how a host responds to SYN, ACK, FIN, NULL, or Xmas packets. UDP scanning required a different understanding because absence of response does not always mean that a port is closed. The project therefore treated scan output as evidence that must be parsed carefully rather than as a simple yes-or-no answer.

The OSI model was also reviewed to build conceptual clarity. Although real-world tools usually operate through TCP/IP terminology, the OSI model is useful for explaining where different controls and attacks occur. Physical and data link layers relate to media access, MAC addresses, and switching. The network layer relates to IP routing and ICMP. The transport layer relates to TCP and UDP ports. The session, presentation, and application layers relate to user-facing protocols and service logic. For this project, the OSI model helped in separating topology discovery from service enumeration and vulnerability analysis. Topology mapping depends heavily on network-layer and transport-layer observations, while CVE matching depends on application-layer product and version evidence.

IP addressing was studied in both practical and security contexts. IPv4 addresses identify hosts in a 32-bit address space, usually written in dotted decimal notation. CIDR notation allows a user to define a network range such as 192.168.1.0/24. Subnetting matters because a scan target may represent one host, a small department subnet, or a larger enterprise segment. The application needed to accept targets such as IP addresses, hostnames, and CIDR ranges, and it needed to group discovered devices into subnet labels for heat maps and topology views. The project also considered IPv6 support because the Nmap profile list includes an IPv6 mode. Even when a demonstration network uses IPv4, a modern scanner should not be designed with assumptions that make IPv6 impossible later.

Ports were studied as logical endpoints through which services communicate. TCP port 22 often indicates SSH, port 80 indicates HTTP, port 443 indicates HTTPS, port 445 indicates SMB, port 3389 indicates RDP, and port 161 indicates SNMP. However, port numbers alone are not sufficient because services may run on non-standard ports. Nmap service detection helps by sending probes and examining banners. In the project, each open port was stored with protocol, state, service, product, version, extra information, tunnel, and CPE data when available. This structure allowed the risk analyzer to consider both the existence of exposure and the nature of the exposed service. A host with many open ports is not automatically critical, but exposed remote administration and known vulnerable versions can significantly increase risk.

Nmap basics were a major part of the internship. Nmap is a widely used network exploration and security auditing tool. It can identify live hosts, discover open ports, detect services, perform OS fingerprinting, run scripts, and output results in several formats. The project used Nmap as the scan engine rather than trying to recreate packet scanning from scratch. This was an important engineering decision because Nmap already includes years of protocol handling, scan options, and fingerprint databases. The application focused on integration, parsing, visualization, and reporting. Basic scan profiles such as quick scanning, ping sweep, service detection, and OS detection were mapped into backend command profiles so that users could run common assessments without remembering every command-line flag.

Advanced Nmap scanning introduced more specialized profiles. A SYN scan is commonly used for efficient TCP discovery, while ACK scans can help reason about firewall behavior. FIN, NULL, and Xmas scans rely on differences in TCP flag handling and may be useful in certain environments, though they are not universally reliable. UDP scans are slower because UDP does not provide the same handshake behavior as TCP. Aggressive scanning combines service detection, OS detection, script scanning, and traceroute-like functionality. The application represented these options as profiles because a graphical tool should guide the user while still preserving flexibility. A custom profile was also included, with output flags stripped so the application could force XML output to standard output and control parsing consistently.

NSE scripts were studied as a powerful extension mechanism in Nmap. The Nmap Scripting Engine allows scripts to perform service enumeration, vulnerability checks, brute-force testing, discovery tasks, and safe informational probes. In the project, NSE scripts were modeled with identifiers, categories, descriptions, average execution estimates, default enablement, and severity. This helped the frontend present script choices without forcing the user to browse a file system. NSE integration also raised ethical and operational questions. Some scripts are safe and informational, while others are intrusive or brute-force oriented. The report therefore treats NSE execution as a controlled assessment capability that should be used only on authorized networks and with awareness of scan impact.

XML parsing was essential because Nmap's XML output is structured and machine-readable. The backend used xml2js to convert XML into JavaScript objects and then normalized host entries into domain models. The parser extracted addresses, hostnames, MAC vendor information, status, response time, OS match, ports, service names, product versions, CPE strings, and script evidence. This step is more than a technical conversion. It is where raw tool output becomes application data. If parsing is weak, the dashboard may mislabel a host, miss a vulnerable service, or generate an incomplete report. The project therefore separated XML parsing from risk analysis so that each responsibility remained clear and testable.

Risk assessment was studied as the process of estimating how serious an exposure is and how urgently it should be addressed. In a production security program, risk depends on asset criticality, exploitability, exposure, control strength, and business impact. In this academic project, a practical scoring model was implemented using open ports, vulnerability severity, and remote administration exposure. The device risk score gives a host-level view, while the attack surface score emphasizes reachable services and high-value administrative ports. These scores do not replace expert judgment, but they help users sort devices and identify which hosts require immediate attention. The dashboard also presents a security score, severity distribution, exposed service lists, and subnet heat maps.

The CVE database concept was studied to understand standardized vulnerability identification. CVE identifiers provide a common name for publicly known cybersecurity vulnerabilities, while NVD and related sources provide enrichment such as descriptions, CVSS scores, affected products, and references. The project included a local CVE mapping database for demonstration and correlation. Product names, versions, operating system families, CVSS values, exploit availability indicators, and patch availability were represented in structured form. During analysis, the system compared service evidence with these records and created vulnerability findings when a match was detected. This approach is simplified compared with enterprise vulnerability scanners, but it demonstrates the core idea behind vulnerability correlation.

OWASP was studied because secure software design is relevant even for a network scanner. A scanner that accepts targets, custom arguments, imported XML, and export requests must defend itself against misuse and malformed input. OWASP categories such as broken access control, cryptographic failures, security misconfiguration, vulnerable components, and identification and authentication failures also appeared in vulnerability findings. The project used OWASP terminology to enrich findings and connect network observations to wider application-security language. This was valuable because real security teams often need to explain how a network finding relates to broader control categories and remediation programs.

React was used for the frontend because the application required an interactive dashboard with several coordinated views. React's component model allowed the scan controls, metric cards, topology graph, device inspector, charts, history panel, report center, settings panel, and live terminal to be developed as separate interface units. State changes from active scans could be reflected immediately in the UI. This was important for user experience because network scanning is not instantaneous; users need to see progress, errors, and results without refreshing the page manually. React also helped organize conditional views such as Overview, Topology, NSE, History, Reports, and Settings.

TypeScript strengthened the project by defining clear data contracts. Types were created for scan profiles, scan phases, severity levels, port services, CVE matches, vulnerability findings, network devices, topology edges, risk metrics, scan summaries, scan records, and application settings. These definitions reduced ambiguity between frontend and backend behavior. For example, a vulnerability finding includes severity, reason, impact, mitigation, references, device IP, port, service, CVE information, MITRE techniques, and OWASP categories. When this structure is typed, it becomes easier to build tables, charts, reports, and filters without guessing field names. TypeScript therefore improved maintainability and reduced runtime mistakes.

Node.js and Express.js were used on the backend to provide APIs around scanning, history, XML analysis, exports, and system status. Node.js was suitable because the project already used TypeScript and because the backend needed to coordinate I/O-heavy work such as spawning Nmap, reading and writing JSON scan history, parsing XML, and streaming scan events. Express.js provided a straightforward HTTP framework for routes such as health checks, system status, NSE script listing, scan creation, scan retrieval, event streaming, XML import, sample scan creation, scan difference comparison, and export generation. Zod validation was used to validate request payloads before they reached scan execution logic.

Report generation was studied as a security communication skill. Raw scan data is useful to a tool, but a human reader needs a different structure: summary, affected hosts, services, vulnerabilities, severity, evidence, and remediation. The project supports HTML, CSV, JSON, and XML export. HTML is useful for readable reports and browser printing. CSV is useful for spreadsheet analysis. JSON preserves structured application data. XML preserves raw Nmap evidence. The report center also supports topology image export in PNG and SVG, which is useful for documentation and presentations. This export capability reflects a real-world need: security findings must move from tools into decision records.

Secure coding practices were applied throughout the project. Input validation limited target length, custom argument length, script selection count, timeout range, and supported scan profiles. The backend avoided shell interpretation when spawning Nmap and constructed commands as argument arrays. Output flags in custom arguments were stripped so the application could keep control over XML output. Imported XML was checked for Nmap structure before analysis. The frontend used typed API calls and error boundaries to reduce failure impact. These practices do not make the project a complete enterprise product, but they show that a cybersecurity tool must itself be built with security awareness.
"""


PROBLEM_STATEMENT = """
Modern organizations operate networks that change continuously. Laptops join and leave wireless networks, servers are deployed for short-term projects, cloud-connected devices expose management ports, and development teams install services that may not be recorded in a central asset register. At the same time, attackers need only one overlooked service, one unpatched component, or one weakly configured remote administration interface to begin reconnaissance and exploitation. This creates a visibility problem before it becomes a vulnerability problem. If administrators do not know which hosts exist, which ports are reachable, which operating systems are present, and which software versions are exposed, they cannot build a reliable remediation plan.

Manual network mapping is difficult because traditional tools present information in fragments. Nmap can provide excellent scan evidence, but its command-line output is not always suitable for quick decision making, especially for students, junior administrators, or teams that need visual summaries. A user may run one command for host discovery, another for service detection, another for operating system detection, and another for NSE vulnerability scripts. The results then need to be copied, interpreted, summarized, and stored. When the network grows, this manual process becomes slow and error-prone. Important observations may be missed because they are buried in text output or spread across different scan files.

Vulnerability assessment also requires context. A service banner may indicate a product and version, but the user still needs to determine whether that version maps to a known CVE, whether the issue is critical or informational, whether a patch exists, and whether the service is exposed on a high-risk port. Without an integrated analyzer, the workflow becomes fragmented across terminal output, web searches, vulnerability databases, spreadsheets, and reports. This fragmentation increases the learning curve and reduces the chance that a scan will lead to timely action. It is especially challenging in academic labs where the objective is to understand the complete security workflow rather than simply run a command.

Another problem is the lack of topology visualization. A list of IP addresses and ports does not show how devices relate to a gateway, which subnets contain higher risk hosts, or which nodes should be inspected first. Visualization is not decoration; it is a way to convert complex technical data into a mental model. A topology graph helps the user see network structure, node relationships, risk concentration, and attack surface in one place. This is valuable for administrators, auditors, and learners because it supports both exploration and explanation.

Finally, many lightweight tools do not provide centralized reporting and history. A scan result is useful at one moment, but vulnerability management depends on comparison over time. Teams need to know whether a host is new, whether a vulnerability has appeared, whether a service was removed, and whether remediation reduced risk. The Network Topology Mapper & Vulnerability Analyzer addresses this problem by combining Nmap scanning, XML parsing, CVE correlation, risk scoring, topology generation, historical storage, and exportable reports into a single web-based platform.
"""


EXISTING_SYSTEM = """
Traditional Nmap usage is powerful but command-oriented. A skilled user can run host discovery, service detection, OS fingerprinting, traceroute, script scans, and custom probes directly from the terminal. This is one of Nmap's strengths because it gives expert users fine control. However, it also creates difficulty for beginners and for report-driven workflows. A user must know which flags to choose, how aggressive a scan should be, how timing templates affect reliability, when UDP scanning is appropriate, and which NSE scripts are safe for a given environment. The output may be human-readable, grepable, XML, or normal text, but the user is still responsible for interpretation and documentation.

General command-line tools share a similar limitation. Utilities such as ping, traceroute, netstat, arp, nslookup, and curl are useful for specific questions, but they do not automatically combine discovery, vulnerability context, topology visualization, and reporting. They are excellent diagnostic instruments, yet they remain separate instruments. When an administrator depends only on these tools, the assessment becomes a sequence of disconnected commands. The final picture must be assembled manually. This approach works for small tasks, but it becomes inefficient when the goal is to understand an entire subnet or produce an internship-style project report with evidence, diagrams, and outcomes.

Wireshark is another important existing system, but its purpose is different. Wireshark captures and analyzes packets. It is extremely useful for protocol troubleshooting, traffic inspection, and forensic analysis. However, it does not automatically scan a network to discover hosts and vulnerabilities. It requires the user to capture traffic at the right point in the network and then interpret packet-level details. For a topology mapper and vulnerability analyzer, Wireshark can support investigation, but it cannot replace active discovery and correlation. It also has a steep learning curve for users who are not already familiar with packet structures and protocol behavior.

Enterprise scanners such as Nessus provide mature vulnerability scanning, plugin ecosystems, reporting, credentialed scanning, and compliance checks. Nessus is a strong professional tool, but it may be expensive or heavy for student projects, small labs, or lightweight internal experiments. It can also feel like a black box for learners because much of the detection logic is hidden inside plugins and proprietary workflows. OpenVAS, now part of the Greenbone ecosystem, is open-source and powerful, but it requires setup, feed management, configuration, and infrastructure that may be more complex than a student needs for a focused topology and Nmap integration project. Both Nessus and OpenVAS are valuable, yet they do not fully match the educational objective of building a custom system that exposes scanning, parsing, risk calculation, visualization, and reporting logic.

Manual documentation is another common existing practice. After running scans, users may paste results into Word documents, spreadsheets, or ticketing tools. This method gives freedom, but it is time-consuming and inconsistent. A port list may be copied without service versions. A CVE reference may be added without evidence. A topology diagram may be drawn manually and become outdated after the next scan. Severity may be assigned differently by different reviewers. When documentation is manual, the report may look complete while still missing important relationships between host inventory, vulnerabilities, risk, and remediation.

The learning curve is therefore a major problem in the existing system. A new user must learn networking, Nmap syntax, scan safety, XML output, vulnerability databases, severity classification, report writing, and visualization separately. Each topic is valuable, but the combined workflow can be intimidating. A fragmented workflow also increases operational risk. If a user forgets to save XML, the parser cannot run. If a user stores results in different folders, historical comparison becomes difficult. If a user does not understand a service banner, a CVE match may be missed. If a user focuses only on critical CVEs, exposed management services without CVE matches may be ignored.

The absence of integrated visualization is especially significant. Many existing workflows end with a table of ports and services. Tables are necessary, but they do not show structure as clearly as a graph. A graph can highlight a gateway, show all discovered devices, visually mark high-risk nodes, and indicate relationship labels such as open-port counts or protocols. Without this view, the user must imagine the network from rows of data. That is possible for a small scan, but it becomes difficult as the number of hosts increases.

The existing system is also time-consuming because each stage depends on manual transfer. Scanning, parsing, CVE research, risk scoring, report creation, and history comparison may require separate tools. Time spent moving data is time not spent understanding risk. The Network Topology Mapper & Vulnerability Analyzer improves this situation by turning the workflow into a single pipeline. The user starts a scan or imports XML, the backend parses the evidence, the analyzer creates findings, the topology engine builds graph data, the dashboard displays results, and the report module exports evidence in multiple formats.
"""


PROPOSED_SYSTEM = """
The proposed system is a web-based cybersecurity platform that integrates network discovery, vulnerability analysis, topology visualization, historical storage, and reporting. Its purpose is not to replace expert judgment, but to reduce the friction between scan execution and security understanding. The application allows a user to enter a target such as an IP address, hostname, or CIDR range, choose a scan profile, optionally select NSE scripts, and start the assessment from a browser interface. The backend executes Nmap locally, collects XML output, parses the result, enriches host and service data, calculates risk scores, builds topology relationships, stores the scan, and returns structured data to the frontend.

The frontend is built using React, TypeScript, Vite, React Flow, Recharts, and a responsive styling layer. It provides a SOC-console style interface with navigation views for Overview, Topology, NSE, History, Reports, and Settings. The Overview view presents metrics such as total hosts, active hosts, critical hosts, open ports, vulnerabilities, unique services, average risk, and security score. Charts summarize severity distribution, exposed services, operating systems, ports, and subnet risk. The Topology view uses React Flow to render an interactive graph where users can zoom, pan, drag nodes, inspect devices, and identify high-risk assets visually. This design solves the no-visualization problem in the existing workflow.

Frontend and backend communication occurs through HTTP APIs and server-sent events. The frontend calls API endpoints to retrieve system status, available NSE scripts, saved scans, sample scans, imported XML analysis, scan exports, and scan differences. When a scan starts, the backend returns a session identifier. The frontend then opens an event stream for that session and receives progress events such as queued, discovery, port scan, service detection, OS detection, NSE, analysis, finished, or error. This live scan progress avoids the uncertainty that users experience when a long command-line scan appears to be stuck. It also creates an audit trail that can be stored with the final scan record.

The scan engine is implemented on the backend using Node.js and Express.js. Scan profiles map user-friendly choices to Nmap arguments. For example, quick scanning uses a faster profile, service detection enables service probing, OS detection enables fingerprinting, aggressive scanning combines several discovery techniques, and custom scanning allows advanced users to provide controlled arguments. Selected NSE scripts are passed through the --script option. The backend forces XML output so the parser receives consistent data. It also validates the request body with Zod, limits selected script count, enforces timeout boundaries, and avoids shell interpretation when spawning the process. These design choices improve safety and reliability.

The parser is responsible for converting Nmap XML into normalized application models. It extracts host addresses, MAC vendor information, hostnames, status, TTL, response time, operating system matches, port records, service names, product versions, extra information, tunnels, and CPE strings. It also collects script evidence and service evidence into text used by the analyzer. The parser assigns subnet labels and infers device type from operating system, vendor, and port patterns. This separation is important because raw XML is too detailed for direct UI use, while a normalized model can support charts, filters, tables, reports, and graph generation.

The risk analyzer evaluates each host after parsing. It uses local rule definitions and CVE records to create vulnerability findings. Rules identify exposures such as anonymous FTP, legacy SMB, weak SSH, legacy TLS, Telnet, exposed RDP, reachable Redis, exposed MongoDB, weak SNMP, and HTTP hardening issues. CVE matching compares product, version, operating system family, and service evidence against local records. Findings include severity, reason, impact, mitigation, references, evidence, affected device, affected port, MITRE technique tags, and OWASP categories. Device risk is then calculated using open port count, severity weight, and remote administration bonuses. Attack surface is calculated using exposed ports, administrative services, and critical findings. The result is a practical prioritization model.

The visualization engine converts analyzed devices into a topology snapshot. Devices are sorted by IP address, a gateway is selected using router/firewall type or common gateway address patterns, and edges connect discovered devices to the gateway. Edge labels show whether the connection represents open ports or a discovered host, and high-risk connections can be animated. The topology graph therefore gives the user a network-oriented view rather than a flat inventory. Filters allow the user to narrow the graph by subnet, vendor, operating system, service, severity, query, or minimum risk score.

The report engine converts scan records into exportable evidence. JSON export preserves the complete application record. XML export preserves the raw Nmap output. CSV export creates rows suitable for spreadsheet analysis, including host, service, risk, vulnerability, and severity fields. HTML export creates a readable report with executive summary, device list, recommendations, and mitigation information. The frontend also supports topology export as PNG or SVG. Together, these outputs solve the centralized-reporting problem by making results portable for academic submission, security review, and future analysis.

The proposed system also includes local JSON scan history. Each completed scan can be saved as a structured file. The History module lists previous scans, allows selection of earlier records, and supports comparison logic for new hosts, removed hosts, new vulnerabilities, and resolved vulnerabilities. This makes the tool useful beyond a one-time scan. It supports continuous learning, remediation tracking, and repeatable assessment. Offline demo mode and XML import provide additional flexibility: the user can replay a sample scan when Nmap is unavailable or import previously collected Nmap XML for analysis.
"""


OBJECTIVES = [
    ("Automated host discovery", "The first objective of the project is to identify active hosts within a user-defined target range. This includes accepting a single host, hostname, or CIDR block and converting scan results into a clear inventory. Automated host discovery reduces the effort required to build a network asset list manually and gives the user a starting point for all later security decisions."),
    ("Network scanning and service enumeration", "The second objective is to scan reachable hosts for open ports, protocols, service names, product banners, and version information. Service enumeration is necessary because risk cannot be understood from host presence alone. A device with SSH, SMB, RDP, HTTP, or database services exposed requires more attention than a device with no reachable services."),
    ("Operating system identification", "The third objective is to infer operating system information wherever scan evidence allows it. OS detection helps the user distinguish between workstations, servers, routers, firewalls, printers, IoT devices, and unknown assets. This classification improves topology interpretation and supports more accurate vulnerability reasoning."),
    ("Interactive topology visualization", "The fourth objective is to present discovered devices as an interactive graph instead of a plain list. The topology view should allow zooming, panning, node selection, minimap navigation, and visual risk highlighting. This objective addresses the practical need for a mental model of the network."),
    ("Vulnerability and CVE correlation", "The fifth objective is to identify potential weaknesses by combining rule-based checks with CVE mapping. The system should examine service names, product versions, OS families, CPE data, and script evidence, then create structured findings with impact and mitigation."),
    ("Risk score and attack surface calculation", "The sixth objective is to convert many technical observations into prioritization indicators. Device risk score, attack surface score, overall network score, security score, and severity distribution help the user decide which hosts require urgent review."),
    ("Historical comparison and reporting", "The seventh objective is to preserve scan history and generate exportable reports. Historical comparison supports remediation tracking, while HTML, CSV, JSON, XML, PNG, and SVG exports support documentation, academic submission, and technical review."),
    ("Safe and maintainable implementation", "The final objective is to build the platform using modern web development practices. React, TypeScript, Express, Zod, NanoID, xml2js, and local JSON storage were selected to create a maintainable codebase with clear interfaces, validation, and structured data flow."),
]


TECH_ROWS = [
    ["React", "Frontend component framework", "Reusable UI, fast state updates, strong ecosystem", "Required for an interactive dashboard with coordinated views"],
    ["TypeScript", "Typed application language", "Compile-time checks, clear contracts, maintainability", "Useful for shared cybersecurity data models"],
    ["Vite", "Frontend build tool", "Fast development server and optimized builds", "Suited to modern React development"],
    ["React Flow", "Topology graph rendering", "Zoom, pan, drag, minimap, edge rendering", "Ideal for interactive network visualization"],
    ["Recharts", "Dashboard charts", "Declarative charts and responsive visual summaries", "Useful for risk, ports, services, and severity charts"],
    ["TailwindCSS", "Utility-first styling approach", "Consistent spacing, responsive layout, rapid UI design", "Selected for a clean SOC-console interface style"],
    ["Node.js", "Backend runtime", "Non-blocking I/O and JavaScript/TypeScript alignment", "Fits scan orchestration, file I/O, and API work"],
    ["Express.js", "HTTP API framework", "Simple routing, middleware support, event endpoints", "Used for scan, history, export, and system APIs"],
    ["Nmap", "Network scan engine", "Mature discovery, port scanning, OS and service detection", "Provides reliable scan evidence without reinventing packet scanning"],
    ["NSE Scripts", "Extended scan checks", "Service enumeration and vulnerability probes", "Adds flexible security assessment capabilities"],
    ["xml2js", "XML parsing", "Converts Nmap XML into JavaScript objects", "Allows deterministic extraction of scan evidence"],
    ["NanoID", "Identifier generation", "Small, URL-safe unique IDs", "Used for scans, findings, sessions, and timeline events"],
    ["Zod", "Request validation", "Runtime schema validation and helpful errors", "Protects scan APIs from malformed inputs"],
    ["Local JSON", "Scan history storage", "Simple, portable, transparent records", "Sufficient for internship scope and offline demonstrations"],
]


ARCH_BLOCKS = [
    ("User", "The user provides scan targets, chooses profiles, selects scripts, imports XML, reviews topology, and exports reports. The user is placed at the top of the architecture because every technical action is initiated by an authorized operator."),
    ("React Frontend", "The frontend converts security data into an interface that can be explored. It contains the dashboard, scan controls, live terminal, topology graph, device table, inspector, NSE manager, history panel, report center, and settings view."),
    ("Express Backend", "The backend exposes APIs, validates requests, starts scan sessions, streams progress, manages history, imports XML, and serves exports. It separates browser interaction from privileged local scan execution."),
    ("Nmap Scan Engine", "The scan engine executes Nmap profiles and selected NSE scripts. It produces XML output for real scans and can also process sample or imported XML when direct scanning is not possible."),
    ("XML Parser", "The parser converts Nmap XML into normalized host, port, service, operating system, vendor, CPE, and script evidence models."),
    ("Risk Analyzer", "The analyzer evaluates open ports, rule findings, CVE matches, severity weights, and administrative exposure to compute device risk and attack surface."),
    ("CVE Correlation Engine", "The CVE component compares service evidence with local CVE records and enriches findings with CVSS, severity, references, patch availability, and exploit indicators."),
    ("Topology Generator", "The topology generator creates nodes and edges, selects a likely gateway, labels relationships, and marks high-risk links for visual attention."),
    ("JSON Database", "Local JSON files store scan records, terminal events, topology snapshots, metrics, summaries, and raw XML for later review."),
    ("Report Generator", "The report generator produces HTML, CSV, JSON, XML, and topology image exports so the assessment can be shared outside the application."),
    ("Dashboard", "The dashboard presents metrics, charts, filters, topology, device details, and findings in a unified view that supports analysis and decision making."),
]


WORKING_STEPS = [
    ("User starts scan", "The workflow begins when the user enters a target and selects a scan profile from the browser interface. The target may be a single IP address, a hostname, or a CIDR range such as 10.10.40.0/24. The user may also select NSE scripts and decide whether to run a real scan, replay a sample, or import XML. This step is important because it translates a technical command-line task into a guided interface. Instead of memorizing scan flags, the user chooses from meaningful options such as quick, service, OS, aggressive, UDP, stealth, traceroute, or custom. The interface also shows system status, Nmap availability, saved scan count, and timeout configuration. By starting from a controlled form, the application reduces accidental misuse and gives the backend a structured request to validate."),
    ("Input validation", "After the scan request reaches the backend, Express and Zod validate the payload. The backend checks that the target is present and within an acceptable length, the scan profile belongs to the supported set, custom arguments are not excessive, the selected script list is bounded, and timeout values remain within a safe range. This validation protects the scan engine from malformed requests and keeps the API behavior predictable. It also supports maintainability because the allowed scan profiles are declared explicitly. Validation is a security requirement because the application interacts with a powerful local tool. A weak backend could allow dangerous argument injection, resource exhaustion, or confusing error states. The project therefore treats validation as part of the scanning workflow rather than as a separate afterthought."),
    ("Nmap execution", "When the request is valid and a real scan is selected, the backend builds an Nmap command using the chosen profile, optional script list, forced XML output, and target. The process is started through a child process without shell interpretation, which reduces command injection risk. The backend monitors stdout and stderr, collects XML output, emits progress events, and applies a timeout so a scan cannot run indefinitely. If Nmap is not installed, the application returns a structured failure record and advises the user to use sample or XML analysis. This design keeps the user informed and prevents silent failure. Nmap remains the authoritative scan engine, while the application controls orchestration, output capture, and result handling."),
    ("XML generation", "Nmap is instructed to produce XML output because XML is structured, machine-readable, and suitable for deterministic parsing. Text output is useful for humans, but XML preserves nested relationships among hosts, addresses, ports, states, services, operating system matches, CPE values, scripts, and timing information. The application strips user-provided output flags in custom mode so that the backend remains in control of where scan output goes. XML output may come from a live scan, an imported file, or a sample fixture. Treating all three sources through the same XML analysis pipeline keeps the architecture consistent and makes offline demonstrations possible."),
    ("XML parsing", "The XML parser uses xml2js to convert the Nmap document into JavaScript objects and then extracts only the fields required by the domain model. It reads host status, IPv4 or IPv6 address, MAC address and vendor, hostname, OS match, response time, TTL, ports, protocols, service names, product names, versions, extra information, tunnels, CPE values, and NSE script text. It normalizes service names and filters invalid host entries. The parser also creates supporting evidence text that later rules and CVE matching can examine. This step is central because every dashboard metric, table row, topology node, and vulnerability finding depends on accurate normalized data."),
    ("Service extraction", "Service extraction focuses on transforming individual port records into useful security observations. A port record includes protocol, port number, state, service name, product, version, extra information, tunnel, and CPE strings. This data allows the application to distinguish between an open port that merely exists and an exposed service that may carry known risk. For example, an open TCP port 22 suggests remote administration, but OpenSSH version evidence can also support CVE matching or weak-configuration rules. Similarly, SMB, RDP, Redis, MongoDB, SNMP, and HTTP services receive special attention because they commonly affect attack surface. Extracted services are later counted for charts, filtered in the UI, included in exports, and used in risk scoring."),
    ("OS detection", "Operating system detection is performed when Nmap provides OS match evidence. The parser records the best available OS string and uses it with vendor and port data to infer a device type such as server, router, firewall, workstation, printer, IoT device, cloud asset, or unknown device. OS detection improves the quality of the topology and report because different device classes require different remediation decisions. An exposed SMB service on a Windows server has a different operational meaning than an unknown embedded device exposing Telnet. OS detection is not perfect, and the report treats it as scan evidence rather than guaranteed truth, but it still improves prioritization and communication."),
    ("Vulnerability analysis", "After parsing, the analyzer evaluates each host against a local rule catalog. The rules check for exposures such as anonymous FTP, legacy SMB, weak SSH indicators, legacy TLS, Telnet, exposed RDP, reachable Redis, exposed MongoDB, weak SNMP, and HTTP hardening gaps. Each rule includes a severity, reason, impact, mitigation, references, MITRE techniques, and OWASP categories. This produces structured vulnerability findings rather than vague warnings. Rule-based analysis is useful because not every risk is represented by a specific CVE. A service can be dangerous due to weak configuration, cleartext communication, or administrative exposure even if the exact product version is unknown."),
    ("CVE mapping", "The CVE mapping engine compares service evidence against local CVE records. It checks product names, version indicators, and operating system families, then creates a CVE-backed finding when the evidence matches. The finding includes CVE ID, CVSS score, severity, description, references, patch availability, exploit availability, matched evidence, impact, and mitigation. This step demonstrates how vulnerability scanners correlate banner and fingerprint data with known vulnerability databases. The implementation is intentionally local for internship scope, but the architecture could later be connected to NVD, CISA KEV, vendor advisories, or threat intelligence feeds. CVE mapping gives findings a standardized identity and makes reports easier to verify."),
    ("Risk score calculation", "Risk score calculation converts detailed findings into a number that supports prioritization. The project calculates device risk by combining open-port exposure, weighted finding severity, and a bonus for remote administration services such as SSH, Telnet, RDP, and SMB. The score is clamped between 0 and 100 and rounded for display. A separate security score is calculated at network level by subtracting weighted risk and attack surface from 100. These scores are not meant to replace professional assessment, but they help the user quickly identify which devices require review. They also support charts, heat maps, filtering, and historical comparison."),
    ("Topology generation", "Topology generation transforms analyzed devices into a graph. Devices are sorted by IP address, a likely gateway is selected using router or firewall classification or common gateway address patterns, and remaining devices are connected to that gateway. Each node contains hostname, IP address, and risk styling. Each edge includes source, target, label, protocol, risk, and animation state. The graph supports zooming, panning, dragging, minimap navigation, and click-based device inspection. This step turns inventory and risk information into a spatial representation that is easier to understand during presentations, audits, and learning exercises."),
    ("Attack surface calculation", "Attack surface calculation estimates how exposed a device is from the perspective of reachable services. The implementation counts open ports, identifies exposed administrative services such as SSH, Telnet, SMB, RDP, VNC, Redis, and MongoDB, and adds weight for critical findings. The resulting score helps distinguish between a host with one low-risk service and a host with multiple management or database services exposed. Attack surface is valuable because attackers often begin with exposure before exploitation. Even when no CVE is detected, a large exposed service set can indicate weak segmentation or unnecessary availability."),
    ("Dashboard visualization", "The dashboard presents scan results through metric cards, charts, filters, tables, topology graph, device inspector, and terminal events. It shows total hosts, active hosts, critical hosts, open ports, vulnerability count, unique service count, average risk, and security score. Recharts visualizations summarize severity distribution, service exposure, operating systems, port frequency, and subnet heat. The device table supports selection and filtering, while the inspector provides host-level detail. This visualization layer makes the analysis usable. Instead of reading raw XML or terminal output, the user can move from summary to device detail and from graph overview to specific evidence."),
    ("Report generation", "Report generation converts the scan record into formats required by different audiences. HTML export provides a readable report that can be opened in a browser or printed. CSV export supports spreadsheet review and sorting. JSON export preserves the complete structured record for future tooling. XML export preserves the original Nmap evidence. The report center can also export topology as PNG or SVG. This step is important because vulnerability management requires communication. A finding that remains inside a dashboard has limited value; a finding that can be shared with a faculty guide, administrator, or remediation team can lead to action."),
    ("History storage", "The final step is storing the scan record in local JSON history. A record includes scan ID, name, target, profile, status, timestamps, command, terminal timeline, topology snapshot, metrics, summary, raw XML, and error information if applicable. Saved history allows the user to revisit earlier scans, compare results, identify new or removed hosts, and track new or resolved vulnerabilities. Local JSON was selected for internship scope because it is transparent, portable, and easy to inspect. In a future enterprise version, this layer could be replaced by a database with authentication, multi-user access, scheduling, and audit logs."),
]


MODULES = [
    ("Authentication (future-ready)", "The current project is designed for local authorized use, but the architecture includes a clear future-ready authentication boundary. In a production version, login, roles, session management, and audit logs would be necessary because scan execution is sensitive. Authentication would prevent unauthorized users from scanning networks, viewing historical results, exporting reports, or changing settings. Role-based access could separate administrators, analysts, read-only auditors, and students. This module is therefore documented as a planned security layer rather than an implemented control."),
    ("Dashboard", "The Dashboard module presents the operational summary of a scan. It includes metric cards, charts, filtering, device tables, topology access, and selected-device inspection. Its purpose is to convert technical evidence into a form that can be understood quickly. A user can see whether a network has critical hosts, which services are most exposed, which operating systems appear most frequently, and which subnet has the highest average risk. The dashboard is the central decision surface of the application."),
    ("Scan Engine", "The Scan Engine module manages Nmap profile construction, selected NSE script inclusion, process execution, timeout handling, XML collection, error handling, and live progress events. It shields the frontend from low-level scan details while still allowing meaningful scan options. The module also supports simulated scans for offline demonstrations. It is the bridge between a web application and a local cybersecurity tool."),
    ("Topology Engine", "The Topology Engine module builds a graph from analyzed host data. It identifies a likely gateway, creates nodes for devices, creates edges for relationships, labels connections, and marks high-risk links. The frontend renders this graph through React Flow with zoom, pan, drag, minimap, and selection behavior. This module makes the project more than a scanner by giving it a network-mapping identity."),
    ("Risk Analyzer", "The Risk Analyzer module calculates device risk, attack surface, overall network score, security score, critical host count, severity distribution, exposed service summaries, OS distribution, port frequency, service frequency, and subnet heat maps. It transforms raw findings into measurable indicators. The module is intentionally transparent so that students and reviewers can understand how scores are produced."),
    ("Vulnerability Engine", "The Vulnerability Engine module combines rule-based findings with CVE correlation. It detects risky services, weak configurations, legacy protocols, and known vulnerable product versions. Each finding includes reason, impact, mitigation, references, evidence, MITRE technique tags, and OWASP categories. This module gives the project its vulnerability assessment capability."),
    ("Report Generator", "The Report Generator module produces exportable outputs in HTML, CSV, JSON, XML, PNG, and SVG forms. It supports both human-readable reports and machine-readable records. The module is important because security work must be documented. Export capability also makes the project useful for academic submission, lab exercises, and administrative review."),
    ("History Module", "The History module saves scan records as local JSON files and allows previous scans to be listed and selected. It also supports comparison between scans by identifying new hosts, removed hosts, new vulnerabilities, and resolved vulnerabilities. This feature moves the application from one-time scanning toward continuous assessment."),
    ("Settings", "The Settings module controls user preferences such as theme, accent color, animation behavior, terminal speed, scan timeout, and auto-refresh. Although these options are not vulnerability controls, they improve usability and allow the tool to adapt to different working styles. Scan timeout is especially practical because network scans can vary significantly in duration."),
]


ALGORITHMS = [
    ("Host Discovery Algorithm", "The host discovery algorithm begins with a target and a selected scan profile. For ping sweep or quick discovery, Nmap is configured to identify responsive hosts without necessarily performing deep service enumeration. The parser then reads host status fields and valid IP addresses. Hosts without usable addresses are ignored, while responsive hosts are normalized with hostname, vendor, response time, TTL, subnet, first-seen time, and last-seen time. The algorithm is practical because it delegates packet-level discovery to Nmap and focuses application logic on consistent interpretation."),
    ("Graph Construction Algorithm", "The graph construction algorithm sorts devices by IP address, chooses a gateway candidate, and creates edges from the gateway to other devices. Router or firewall classification is preferred; otherwise a common gateway pattern such as an address ending in .1 is used; if neither exists, the first sorted host becomes the root. Each edge includes label, protocol, highest observed risk, and animation state. This creates a readable star-oriented topology suitable for subnet-level assessment and academic demonstration."),
    ("Risk Score Calculation Algorithm", "The risk score algorithm counts open ports, adds weighted severity from vulnerability findings, and applies an additional score when remote administration services are exposed. Severity weights allow critical findings to influence the score more than informational findings. The final value is clamped to a 0 to 100 range. This algorithm is simple enough to explain but still meaningful because it combines exposure and vulnerability evidence."),
    ("Attack Surface Analysis Algorithm", "The attack surface algorithm evaluates how reachable and attractive a host appears from an attacker's perspective. It counts open ports, identifies exposed administrative and database services, and adds weight for critical findings. This score is separate from vulnerability risk because a host may be highly exposed even before a specific CVE match is found. The algorithm helps users find systems that need segmentation, firewall review, or service reduction."),
    ("CVE Matching Algorithm", "The CVE matching algorithm creates a lowercase evidence string from service name, product, version, extra information, tunnel, operating system, and CPE values. It then checks local CVE records for product matches, version matches, and operating system family matches. When a record matches, the analyzer creates a structured vulnerability finding with CVE ID, CVSS score, severity, description, reference links, patch status, exploit status, and matched evidence. The algorithm demonstrates the core logic of vulnerability correlation."),
    ("Network Relationship Detection Algorithm", "The relationship detection algorithm uses topology assumptions and scan evidence to describe how nodes should be connected in the graph. In this version, all hosts are connected to a likely gateway because the tool focuses on subnet-level visualization. Edge labels use open-port counts when available and host-discovered labels otherwise. Risk coloring is derived from related vulnerabilities. Future versions could improve this algorithm using traceroute hops, ARP data, switch information, VLAN data, and routing tables."),
]


FEATURES = [
    ("Host Discovery", "Host discovery identifies which systems respond within the selected target range. The feature is implemented through Nmap profiles and XML parsing, then displayed as an inventory of network devices. Each discovered host is assigned an IP address, hostname when available, vendor, subnet, status, first-seen time, and last-seen time. This feature is essential because all later analysis depends on knowing which devices exist. It reduces manual asset discovery and gives the user an immediate view of active systems."),
    ("Network Scanning", "Network scanning extends host discovery by examining ports, protocols, service states, and scan profiles. The application supports quick, intense, service, version, OS, UDP, SYN, ACK, NULL, FIN, Xmas, stealth, aggressive, IPv6, ping sweep, traceroute, and custom modes. This flexibility allows the user to choose between speed, depth, and specialized behavior. Scanning is controlled from the frontend but executed by the backend, keeping privileged tool interaction away from the browser."),
    ("Service Detection", "Service detection identifies applications listening behind open ports. The parser records service name, product, version, extra information, tunnel, and CPE values when Nmap provides them. This evidence allows the tool to distinguish between generic exposure and specific vulnerable products. Service detection supports dashboard charts, device inspection, CSV export, CVE matching, and rule-based vulnerability analysis. Without service detection, the tool would know that ports are open but not what security meaning they carry."),
    ("OS Detection", "OS detection records Nmap operating system matches and uses them to improve device classification. Knowing whether a host appears to be Windows, Linux, network infrastructure, or an unknown embedded system helps users interpret risk. OS detection also supports CVE matching because some vulnerabilities apply only to specific platform families. The project treats OS detection as best-effort evidence and does not assume perfect accuracy. This cautious approach is important because fingerprinting can be affected by firewalls, NAT, and limited probe responses."),
    ("Interactive Network Topology", "The interactive topology feature presents scan results as a graph. Users can zoom, pan, drag nodes, view a minimap, and click devices for inspection. Nodes are styled by risk band, and high-risk edges can be visually emphasized. This feature changes the user experience from reading lists to exploring a network model. It is especially valuable in presentations because a topology graph communicates structure and risk faster than a long table."),
    ("Node Relationship Mapping", "Node relationship mapping converts device records into graph edges. The current system selects a likely gateway and maps other devices to it, labeling relationships with open-port counts or host-discovered evidence. While simplified, this approach is effective for subnet-level visualization. It gives users a clear representation of which assets are part of the scanned segment and which nodes require deeper review. Future versions can enrich relationships with traceroute, ARP, VLAN, and routing data."),
    ("Risk Score Calculation", "Risk score calculation produces a 0 to 100 score for each device. It considers open services, vulnerability severity, and exposed remote administration ports. The score helps users sort devices and identify high-priority systems. A high score does not automatically prove compromise, but it signals that the device deserves attention. Risk scoring is useful in academic and operational contexts because it teaches prioritization rather than treating every finding as equal."),
    ("Attack Surface Analysis", "Attack surface analysis estimates how exposed a device is. It counts open ports, highlights administrative services, and considers critical findings. This feature is different from CVE detection because exposure can be risky even without a known vulnerability. For example, public RDP, Telnet, SMB, Redis, or MongoDB may require segmentation and access control regardless of exact version. The feature teaches users to think about reachable services as part of defensive design."),
    ("CVE Detection", "CVE detection maps service evidence to known vulnerability identifiers. The project includes local CVE records for demonstration and correlation. When product, version, and operating system evidence match a record, the analyzer creates a finding with CVE ID, CVSS score, severity, references, patch availability, exploit availability, and mitigation. This feature connects network scanning with vulnerability intelligence and helps users produce reports that reference standardized identifiers."),
    ("NSE Script Execution", "NSE script execution allows the user to select scripts that extend Nmap's behavior. Scripts may perform discovery, service enumeration, vulnerability checks, authentication review, HTTP inspection, SMB analysis, SSH algorithm enumeration, DNS checks, SNMP review, and more. The application models scripts with categories, descriptions, risk levels, and default enablement. This helps users choose scripts with better context. NSE support makes the platform more flexible than a fixed-port scanner."),
    ("Vulnerability Severity Classification", "Severity classification organizes findings into critical, high, medium, low, and informational bands. Each finding receives severity from either the local rule catalog or CVE data. Severity is used in dashboards, filters, topology colors, reports, and risk calculations. Classification makes output easier to act upon because users can focus first on critical remote code execution, exposed administrative access, or severe misconfiguration before reviewing lower-priority observations."),
    ("Historical Scan Comparison", "Historical scan comparison allows the application to compare two saved scan records. It identifies new hosts, removed hosts, new vulnerabilities, and resolved vulnerabilities. This feature supports remediation tracking and change awareness. In a real environment, risk is dynamic; new services appear and old vulnerabilities may return. Historical comparison teaches that scanning should be a repeated process rather than a one-time activity."),
    ("Live Scan Progress", "Live scan progress is implemented through server-sent events. While Nmap runs, the backend emits timeline events for phases such as queued, discovery, port scan, service detection, NSE, analysis, finished, and error. The frontend displays these events in a live terminal. This feature improves usability because long scans can otherwise feel uncertain. It also creates an activity record that is saved with the scan and can be reviewed later."),
    ("Report Export", "Report export converts scan data into shareable evidence. The project supports multiple formats so that different users can consume the same assessment in different ways. A faculty reviewer may prefer a readable HTML report, an analyst may prefer CSV, a developer may prefer JSON, and a security engineer may want raw XML. Export support also demonstrates that cybersecurity tools must communicate results, not merely collect them."),
    ("HTML Export", "HTML export creates a browser-readable report with summary, device list, findings, and recommendations. It can be opened directly, printed, or converted to PDF using browser tools. HTML is useful because it preserves structure and readability without requiring specialized software. In the project, this format is designed for quick review and academic presentation. It gives the user a clean report outside the live dashboard."),
    ("CSV Export", "CSV export creates spreadsheet-friendly rows for hosts, services, risks, vulnerabilities, and severities. This format is useful for filtering, sorting, pivoting, and manual remediation tracking. Many organizations still rely on spreadsheets for short-term vulnerability management, so CSV remains valuable even when dashboards exist. The export makes it easier to hand findings to another reviewer or combine scan output with other records."),
    ("JSON Export", "JSON export preserves the full structured scan record. It includes topology, devices, ports, findings, metrics, summary, terminal events, and raw XML when available. JSON is useful for developers, automation, and future integrations. It also makes the project transparent because the same data shown in the UI can be inspected directly. This format supports extensibility and debugging."),
    ("XML Export", "XML export returns the original Nmap XML output when available. This is important for evidence preservation because XML is the source from which parsed records were derived. A security reviewer can reprocess the XML with other tools, verify parsing assumptions, or archive the raw scan result. XML export therefore supports traceability between the scanner, analyzer, and report."),
    ("Offline Demo Mode", "Offline demo mode allows the system to replay a sample Nmap XML scan when Nmap is not installed or a live network cannot be scanned. This is useful in academic demonstrations, classrooms, and restricted environments. The feature ensures that the dashboard, parser, analyzer, topology graph, history, and report modules can still be demonstrated. It also separates learning the application workflow from the practical limitations of scan permissions and local tool installation."),
]


ADVANTAGES = [
    ("Reduced manual effort", "The application reduces the manual effort required to move from scanning to understanding. Instead of running multiple commands, saving separate files, copying output into a document, searching for CVEs, and drawing diagrams manually, the user can run or import a scan and receive normalized results in one interface. This saves time and reduces transcription errors."),
    ("Centralized network visibility", "The platform centralizes host inventory, services, vulnerabilities, risk metrics, topology, history, and exports. Centralization is valuable because security decisions depend on seeing relationships among assets, exposure, and findings. A centralized view also helps students and administrators explain results to others without switching between many tools."),
    ("Improved prioritization", "Risk score, attack surface score, severity distribution, and critical-host count support prioritization. Not every open port requires the same urgency. The application helps users identify hosts with critical findings, exposed remote administration, or large service exposure. This encourages a remediation mindset rather than a simple inventory mindset."),
    ("Better learning experience", "The project is suitable for cybersecurity education because it exposes each stage of the assessment workflow. Students can see how Nmap output becomes parsed records, how services become findings, how findings become risk scores, and how records become reports. This transparency is useful for learning because it avoids treating vulnerability scanning as a black box."),
    ("Multiple export options", "HTML, CSV, JSON, XML, PNG, and SVG exports make the tool flexible. Different formats serve different purposes: readable reports, spreadsheets, automation, raw evidence, and topology diagrams. Export diversity improves the practical usefulness of the project and supports academic submission requirements."),
    ("Open-source oriented stack", "The project uses widely available technologies such as React, TypeScript, Node.js, Express, Nmap, xml2js, NanoID, and Zod. This keeps cost low and makes the architecture understandable. Students can inspect and extend the code rather than depending on expensive proprietary scanning platforms."),
    ("Cross-platform potential", "Because the application is based on Node.js, a browser frontend, and Nmap, it can run on Windows, Linux, and macOS with appropriate setup. Nmap availability and privileges differ by platform, but the general architecture remains portable. This matters for student and lab environments where machines vary."),
    ("Extensible architecture", "The system is organized into frontend components, API routes, scan profiles, parser logic, rule catalog, CVE records, analysis engine, history service, and report service. This modular structure makes future enhancement easier. New scan profiles, rules, CVE sources, export formats, authentication, or database backends can be added without rewriting the entire application."),
]


LIMITATIONS = [
    ("Requires Nmap installation", "Real scanning depends on the Nmap executable being installed on the host system. If Nmap is missing, the application can still analyze sample or imported XML, but it cannot perform live discovery. This limitation is expected because the project intentionally uses Nmap as the trusted scan engine instead of reimplementing scanner logic."),
    ("Administrative privileges", "Some scan types require elevated privileges, especially SYN scanning, OS detection, certain UDP behavior, and low-level packet techniques. Without administrative rights, Nmap may fall back to less precise methods or fail. Users must therefore understand platform permissions and run scans within authorized environments."),
    ("Firewall restrictions", "Firewalls, host-based security tools, NAT, routing rules, and intrusion prevention systems can change scan visibility. A filtered port may represent a firewall decision rather than a closed service. A host may appear down because it blocks probes. The application reports evidence but cannot always determine hidden network truth."),
    ("False positives", "Rule-based findings and CVE matches can produce false positives when service banners are misleading, versions are backported, or evidence is incomplete. The report should be treated as an assessment aid that requires validation. This limitation is common in vulnerability scanning and should be communicated honestly."),
    ("Scan duration", "Deep scans, UDP scans, full-port scans, NSE scripts, and large target ranges can take significant time. The application includes live progress and timeout settings, but it cannot remove the underlying cost of network probing. Users must choose profiles according to network size and acceptable scan impact."),
    ("Large networks", "The current topology model and local JSON storage are suitable for small to medium lab networks. Very large enterprise networks may require pagination, database indexing, distributed scanning, credential management, asset ownership data, and more advanced graph layout. The project is designed as an internship platform rather than an enterprise scanner."),
    ("Offline database limitations", "The local CVE mapping database is useful for demonstration but cannot match the completeness or freshness of live vulnerability feeds. New CVEs appear frequently, and product matching is complex. A future version should integrate NVD, CISA KEV, vendor advisories, and threat intelligence sources for better coverage."),
]


FUTURE_SCOPE = [
    ("AI-based vulnerability prioritization", "A future version could use intelligent prioritization models to combine CVSS, exploit availability, asset role, exposure, business criticality, and historical activity into richer remediation guidance. This would help analysts focus on issues most likely to matter in their environment."),
    ("MITRE ATT&CK integration", "The project already stores MITRE technique tags in findings. Future work can expand this into a full ATT&CK mapping view that explains possible attacker tactics, techniques, and defensive controls associated with each exposure."),
    ("Shodan API integration", "Shodan integration could help compare internal scan findings with internet-exposed asset intelligence. This would be useful for organizations that need to know whether a service is visible beyond the local network."),
    ("CISA KEV integration", "The CISA Known Exploited Vulnerabilities catalog could be integrated to highlight vulnerabilities that are known to be exploited in the wild. KEV status would make prioritization more realistic than relying on CVSS alone."),
    ("Real-time monitoring", "The platform can evolve from on-demand scanning to scheduled or continuous monitoring. Real-time or periodic checks would identify new hosts and services soon after they appear, improving response time."),
    ("Cloud asset discovery", "Future versions can discover cloud assets through provider APIs such as AWS, Azure, and Google Cloud. This would expand the tool beyond local networks and support hybrid environments."),
    ("Docker deployment", "Containerized deployment would simplify installation and make the backend, frontend, and dependencies easier to run consistently. Docker support could also help instructors deploy the project in lab environments."),
    ("Authentication", "User login, session management, role-based access, and audit logging should be added before production use. Authentication is important because scan results and scan execution capabilities are sensitive."),
    ("Multi-user support", "A database-backed multi-user model could allow teams to share scans, assign remediation owners, comment on findings, and maintain separate permissions for analysts, administrators, and viewers."),
    ("Email alerts", "Email notifications could inform users when scans finish, when critical vulnerabilities are detected, or when new hosts appear. Alerts would make the system more proactive."),
    ("PDF reports", "Direct PDF report generation would improve formal reporting. Although HTML can be printed, a dedicated PDF engine could create consistent formatting, page numbers, cover pages, and executive summaries."),
    ("Machine learning", "Machine learning could be used to classify devices, detect unusual service combinations, cluster similar hosts, and identify changes that deviate from normal network behavior."),
    ("Threat intelligence", "Threat intelligence feeds could enrich CVE findings with exploit activity, malware associations, attacker interest, and recommended urgency. This would connect local evidence to external threat context."),
    ("Zero Trust", "Zero Trust concepts could be added through identity-aware access recommendations, segmentation guidance, least-privilege service exposure, and continuous verification checks."),
    ("Continuous scanning", "Continuous scanning would allow the tool to maintain an updated view of the network. Combined with history comparison, it could show whether remediation is improving security posture over time."),
]


TEST_TABLES = {
    "Functional Testing": [
        ["Start scan with valid target", "User enters 10.10.40.0/24 and selects aggressive scan", "Session is created and progress events begin", "Pass"],
        ["Import XML", "User uploads Nmap XML file", "Hosts, ports, services, and findings are parsed", "Pass"],
        ["Load demo", "User clicks sample/demo option", "Sample scan appears in dashboard", "Pass"],
        ["Export report", "User chooses HTML, CSV, JSON, or XML", "Correct content type and file content are returned", "Pass"],
        ["Filter devices", "User filters by risk, service, or severity", "Topology and table update consistently", "Pass"],
    ],
    "Performance Testing": [
        ["Small subnet scan", "Sample /24 XML fixture", "Dashboard loads without noticeable delay", "Pass"],
        ["Large scan record", "Many hosts with multiple ports", "JSON parsing and chart rendering remain usable", "Acceptable"],
        ["Event stream", "Long-running scan progress", "Terminal receives bounded event history", "Pass"],
        ["Export generation", "Scan with many findings", "CSV and HTML exports are generated quickly", "Pass"],
    ],
    "Security Testing": [
        ["Invalid profile", "Send unsupported scan profile", "API rejects request through schema validation", "Pass"],
        ["Oversized scripts", "Submit more scripts than allowed", "API rejects or truncates according to schema", "Pass"],
        ["Custom output flags", "Provide custom -oX or -oN output flags", "Backend strips conflicting output flags", "Pass"],
        ["Missing Nmap", "Run system without Nmap executable", "Structured error record is generated", "Pass"],
    ],
    "Compatibility Testing": [
        ["Modern browser", "Open dashboard in Chromium-based browser", "UI renders and API calls work", "Pass"],
        ["Localhost backend", "Run Express API on local port", "Frontend communicates with backend", "Pass"],
        ["XML fixture", "Analyze saved XML independent of live Nmap", "Output matches expected domain model", "Pass"],
        ["Cross-platform path behavior", "Use Node.js file APIs for history", "History storage remains portable", "Pass"],
    ],
    "Usability Testing": [
        ["Navigation", "Switch between Overview, Topology, NSE, History, Reports, Settings", "Views are clear and responsive", "Pass"],
        ["Scan control clarity", "Select profile and scripts", "Options are visible and understandable", "Pass"],
        ["Risk interpretation", "Review high-risk node", "Device inspector shows evidence and findings", "Pass"],
        ["Report access", "Use report center", "Export options are grouped logically", "Pass"],
    ],
}


BIBLIOGRAPHY = [
    "[1] Nmap Project, \"Nmap Reference Guide,\" Nmap.org. Available: https://nmap.org/book/man.html",
    "[2] Nmap Project, \"Nmap Scripting Engine Documentation,\" Nmap.org. Available: https://nmap.org/nsedoc/",
    "[3] OWASP Foundation, \"OWASP Top 10:2021,\" OWASP.org. Available: https://owasp.org/Top10/",
    "[4] MITRE, \"MITRE ATT&CK Enterprise Matrix,\" attack.mitre.org. Available: https://attack.mitre.org/",
    "[5] Cybersecurity and Infrastructure Security Agency, \"Known Exploited Vulnerabilities Catalog,\" cisa.gov. Available: https://www.cisa.gov/known-exploited-vulnerabilities-catalog",
    "[6] National Institute of Standards and Technology, \"National Vulnerability Database,\" nvd.nist.gov. Available: https://nvd.nist.gov/",
    "[7] CVE Program, \"Common Vulnerabilities and Exposures,\" cve.org. Available: https://www.cve.org/",
    "[8] React Team, \"React Documentation,\" react.dev. Available: https://react.dev/",
    "[9] Express.js, \"Express Web Framework Documentation,\" expressjs.com. Available: https://expressjs.com/",
    "[10] Node.js Project, \"Node.js Documentation,\" nodejs.org. Available: https://nodejs.org/en/docs/",
    "[11] TypeScript Team, \"TypeScript Documentation,\" typescriptlang.org. Available: https://www.typescriptlang.org/docs/",
    "[12] Vite Team, \"Vite Guide,\" vite.dev. Available: https://vite.dev/guide/",
    "[13] Internet Engineering Task Force, \"RFC 791: Internet Protocol,\" Sept. 1981. Available: https://www.rfc-editor.org/rfc/rfc791",
    "[14] Internet Engineering Task Force, \"RFC 793: Transmission Control Protocol,\" Sept. 1981. Available: https://www.rfc-editor.org/rfc/rfc793",
    "[15] Internet Engineering Task Force, \"RFC 768: User Datagram Protocol,\" Aug. 1980. Available: https://www.rfc-editor.org/rfc/rfc768",
]


def add_cover(doc: Document) -> None:
    add_para(doc, "LOVELY PROFESSIONAL UNIVERSITY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18, color=BLACK, after=18)
    add_para(doc, "SUMMER TRAINING / INTERNSHIP REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, color=BLUE, after=28)
    add_para(doc, "Project Title:", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, after=4)
    add_para(doc, "NETWORK TOPOLOGY MAPPER & VULNERABILITY ANALYZER", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=20, color=DARK_BLUE, after=34)
    add_para(doc, "Submitted By: ______________________________", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=8)
    add_para(doc, "Registration No.: ___________________________", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=8)
    add_para(doc, "School of Computer Science and Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=8)
    add_para(doc, "Term: Aug-Dec 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=44)
    add_para(doc, "Submitted in partial fulfillment of the requirements for Summer Training / Internship", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=11, color=MUTED, after=8)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    doc.add_heading("DECLARATION", level=1)
    add_body(doc, "I hereby declare that this internship report titled 'Network Topology Mapper & Vulnerability Analyzer' is my original work carried out during the internship/training period. The work presented in this report has been prepared for academic submission at Lovely Professional University and has not been submitted previously for any degree or diploma. The project, analysis, documentation, and conclusions have been developed with proper academic care, and all external references used for technical understanding have been acknowledged in the bibliography.")
    add_para(doc, "Signature of Student: ______________________________", after=4)
    add_para(doc, "Date: ____________________", after=12)
    doc.add_page_break()

    doc.add_heading("CERTIFICATE", level=1)
    add_body(doc, "This is to certify that the project titled 'Network Topology Mapper & Vulnerability Analyzer' was successfully completed as part of the Summer Training / Internship requirements of the School of Computer Science and Engineering, Lovely Professional University. The project demonstrates practical understanding of network discovery, Nmap integration, vulnerability analysis, CVE correlation, topology visualization, historical scan management, and report generation. To the best of our knowledge, the work presented in this report is original and has been carried out with sincere effort during the internship period.")
    add_para(doc, "Faculty Guide: ______________________________        Student: ______________________________", after=4)
    add_para(doc, "Name: ______________________________                 Registration No.: ______________________", after=12)
    doc.add_page_break()

    doc.add_heading("ACKNOWLEDGEMENT", level=1)
    add_body(doc, "I express my sincere gratitude to Lovely Professional University and the School of Computer Science and Engineering for providing me with the opportunity to undertake this Summer Training / Internship project. The project allowed me to connect classroom concepts with practical implementation in networking, cybersecurity, software engineering, and technical documentation. I am thankful to my faculty mentor for continuous guidance, valuable suggestions, and encouragement throughout the work. I also acknowledge the support of my peers, friends, and family, whose motivation helped me complete the project with discipline and confidence. This report is the result of consistent learning, experimentation, debugging, and documentation, and I sincerely thank everyone who contributed directly or indirectly to its completion.")
    add_para(doc, "Student Signature: ______________________________", after=12)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_heading("TABLE OF CONTENTS", level=1)
    entries = [
        "1. Cover Page",
        "2. Declaration",
        "3. Certificate",
        "4. Acknowledgement",
        "5. Table of Contents",
        "6. Introduction of Organization",
        "7. Internship Content",
        "8. Project Details",
        "8.1 Problem Statement",
        "8.2 Objectives",
        "8.3 Existing System",
        "8.4 Proposed System",
        "8.5 Technologies Used",
        "8.6 System Architecture",
        "8.7 Working of the Project",
        "8.8 Modules",
        "8.9 Algorithms",
        "8.10 Key Features",
        "8.11 Advantages",
        "8.12 Limitations",
        "8.13 Future Scope",
        "9. Testing",
        "10. Results",
        "11. Conclusion",
        "12. Bibliography",
    ]
    for entry in entries:
        add_para(doc, entry, after=2)
    doc.add_page_break()


def add_project_details(doc: Document) -> None:
    doc.add_heading("8. PROJECT DETAILS", level=1)
    doc.add_heading("8.1 Problem Statement", level=2)
    add_body(doc, PROBLEM_STATEMENT)

    doc.add_heading("8.2 Objectives", level=2)
    for heading, text in OBJECTIVES:
        doc.add_heading(heading, level=3)
        add_body(doc, text)

    doc.add_heading("8.3 Existing System", level=2)
    add_body(doc, EXISTING_SYSTEM)

    doc.add_heading("8.4 Proposed System", level=2)
    add_body(doc, PROPOSED_SYSTEM)

    doc.add_heading("8.5 Technologies Used", level=2)
    add_table(doc, ["Technology", "Purpose", "Advantages", "Reason for Selection"], TECH_ROWS, [1.25, 1.55, 1.85, 1.85])
    for row in TECH_ROWS:
        doc.add_heading(row[0], level=3)
        add_body(doc, f"{row[0]} is used for {row[1].lower()}. Its main advantages in this project are {row[2].lower()}. It was selected because {row[3].lower()}. In the overall architecture, {row[0]} contributes to a controlled, maintainable, and reportable security workflow rather than existing as an isolated technology choice.")

    doc.add_heading("8.6 System Architecture", level=2)
    add_ascii_block(doc, """
User
  |
  v
React Frontend
  |
  v
Express Backend
  |
  v
Nmap Scan Engine
  |
  v
XML Parser
  |
  v
Risk Analyzer
  |
  v
CVE Correlation Engine
  |
  v
Topology Generator
  |
  v
JSON Database
  |
  v
Report Generator
  |
  v
Dashboard
""")
    for heading, text in ARCH_BLOCKS:
        doc.add_heading(heading, level=3)
        add_body(doc, text)

    doc.add_heading("8.7 Working of the Project", level=2)
    for index, (heading, text) in enumerate(WORKING_STEPS, 1):
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(f"{heading}. ")
        set_run_font(run, bold=True)
        run = paragraph.add_run(text)
        set_run_font(run)

    doc.add_heading("8.8 Modules", level=2)
    for heading, text in MODULES:
        doc.add_heading(heading, level=3)
        add_body(doc, text)

    doc.add_heading("8.9 Algorithms", level=2)
    for heading, text in ALGORITHMS:
        doc.add_heading(heading, level=3)
        add_body(doc, text)

    doc.add_heading("8.10 Key Features", level=2)
    for heading, text in FEATURES:
        doc.add_heading(heading, level=3)
        add_body(doc, text)

    doc.add_heading("8.11 Advantages", level=2)
    for heading, text in ADVANTAGES:
        doc.add_heading(heading, level=3)
        add_body(doc, text)

    doc.add_heading("8.12 Limitations", level=2)
    for heading, text in LIMITATIONS:
        doc.add_heading(heading, level=3)
        add_body(doc, text)

    doc.add_heading("8.13 Future Scope", level=2)
    for heading, text in FUTURE_SCOPE:
        doc.add_heading(heading, level=3)
        add_body(doc, text)


def add_testing(doc: Document) -> None:
    doc.add_heading("9. TESTING", level=1)
    add_body(doc, "Testing was performed to verify that the application behaves correctly as a network assessment platform and as a web application. The testing approach covered functional behavior, performance expectations, security controls, compatibility, and usability. Since the project integrates a local system tool with a browser interface, testing focused not only on successful output but also on safe handling of invalid inputs, missing dependencies, imported XML, scan history, and report exports.")
    for heading, rows in TEST_TABLES.items():
        doc.add_heading(heading, level=2)
        add_body(doc, f"{heading} was used to confirm that the corresponding area of the project satisfied expected behavior under realistic internship-lab conditions.")
        add_table(doc, ["Test Case", "Input / Condition", "Expected Result", "Status"], rows, [1.55, 1.85, 2.35, 0.75])
    add_body(doc, "The testing results show that the major workflow is stable for demonstration and academic submission. Live scans depend on local Nmap availability and permissions, but imported XML and demo mode allow the analyzer, dashboard, topology, history, and export functions to be tested even when live scanning is not possible. The most important observation from testing is that validation and structured error handling improve trust in the tool. A scanner should not fail silently or return unclear output; it should explain whether a problem occurred during input validation, Nmap execution, XML parsing, or analysis.")


def add_results(doc: Document) -> None:
    doc.add_heading("10. RESULTS", level=1)
    add_body(doc, """
The project successfully produced a working cybersecurity platform that combines network discovery, topology mapping, vulnerability analysis, risk scoring, historical storage, live progress, and report export. The obtained outputs demonstrate that Nmap XML can be transformed into a richer security record. Instead of leaving the user with terminal text, the application produces devices, ports, services, vulnerabilities, topology edges, metrics, timeline events, and exportable evidence. This confirms that the project met its main objective of integrating scanning with visualization and reporting.

Risk scores were generated for discovered devices by considering open ports, vulnerability severity, and exposed administration services. The dashboard displayed average risk, security score, critical host count, vulnerability count, and attack surface indicators. These metrics helped identify which devices required immediate inspection. The most useful result was not merely the number itself but the combination of score, evidence, and mitigation. A user could select a high-risk node, inspect open ports, review findings, and understand why the device received that score.

Topology visualization was obtained through an interactive graph. Devices appeared as nodes with hostnames and IP addresses, while relationships were drawn as edges from a likely gateway to discovered hosts. Risk-based coloring and animation helped distinguish high-priority assets. The topology view was particularly effective for explaining scan results because it gave a visual structure to the network. The user could zoom, pan, drag, and inspect nodes, making the output suitable for classroom demonstration and technical review.

Scan accuracy depended on the quality of Nmap evidence. For services with clear product and version banners, the parser and analyzer produced useful findings. For filtered hosts, ambiguous banners, or missing OS evidence, the application appropriately preserved uncertainty by using unknown values or evidence-based classification. This reflects realistic cybersecurity practice: scan results provide evidence, not omniscience. The system improved accuracy by separating raw XML, normalized data, rule findings, CVE matches, and calculated metrics, allowing each layer to be reviewed independently.

Performance was acceptable for the intended internship scope. The frontend remained responsive with sample scan data, the backend generated records from XML, and exports were produced without complex infrastructure. Local JSON history was sufficient for small and medium demonstrations. Larger enterprise networks would require database optimization and more advanced graph layout, but the current implementation is appropriate for an academic project. Overall, the result is a complete prototype that demonstrates the end-to-end workflow of scanning, analyzing, visualizing, documenting, and preserving network security evidence.
""")


def add_conclusion(doc: Document) -> None:
    doc.add_heading("11. CONCLUSION", level=1)
    add_body(doc, """
The Network Topology Mapper & Vulnerability Analyzer project successfully demonstrates how cybersecurity concepts and modern web development can be combined to solve a practical network visibility problem. The project began with a simple but important observation: many users can run a scanner, but fewer can convert scan output into a clear understanding of network structure, exposed services, vulnerability severity, risk priority, and reportable evidence. By integrating Nmap, NSE scripts, XML parsing, CVE mapping, rule-based analysis, topology visualization, local history, live progress, and multiple export formats, the application turns a fragmented workflow into a structured assessment pipeline.

The project also shows the importance of asset discovery as the first step in security. A vulnerability cannot be managed if the affected host is unknown. A network cannot be defended effectively if administrators do not know which services are reachable. The application therefore treats host discovery, service detection, OS fingerprinting, and topology construction as foundational security activities. Once the network picture is available, the analyzer adds meaning through severity classification, CVE correlation, attack surface calculation, and risk scoring. This sequence reflects real defensive practice: discover, understand, prioritize, report, and improve.

From a software engineering perspective, the project confirms the value of a modular architecture. The React frontend focuses on presentation and interaction. The Express backend handles validation, scan orchestration, event streaming, XML import, history, and exports. The parser converts Nmap XML into normalized data. The risk engine creates findings and metrics. The topology generator builds graph relationships. The report service transforms records into portable formats. Because these responsibilities are separated, the system is easier to maintain and extend. Future features such as authentication, database storage, CISA KEV integration, scheduled scans, cloud asset discovery, PDF reporting, and multi-user support can be added without discarding the core design.

The project provided significant learning in networking and cybersecurity. Concepts such as TCP/IP, OSI layers, ports, protocols, service banners, CPE data, CVEs, CVSS, OWASP categories, MITRE techniques, and scan profiles became practical rather than theoretical. The implementation clarified why scanning must be authorized, why results require validation, why false positives occur, and why risk depends on exposure as well as vulnerability identity. It also reinforced that cybersecurity tools must be secure themselves. Input validation, bounded script selection, controlled command construction, timeout handling, structured errors, and evidence preservation are necessary in a tool that interacts with local scanning capabilities.

The report also highlights the communication dimension of cybersecurity. A technical finding has limited value if it cannot be explained. The project therefore includes device inspection, charts, topology diagrams, severity labels, mitigation text, and export formats. These outputs support different audiences. A technical reviewer may inspect XML or JSON, a manager may read an HTML summary, a student may study the topology graph, and a remediation owner may use CSV rows. This reinforces an important professional lesson: security assessment is not complete until findings are documented in a way that supports action.

The current system has limitations, including dependency on Nmap installation, possible need for administrative privileges, firewall effects, false positives, scan duration, scalability limits, and local CVE database coverage. However, these limitations do not reduce the academic value of the project. Instead, they show that the implementation is realistic. Real scanners also face incomplete evidence, filtered networks, changing assets, and evolving vulnerability intelligence. A good security tool should make these constraints visible and provide a path for improvement.

In conclusion, the Network Topology Mapper & Vulnerability Analyzer fulfills the objectives of an internship project by combining theory, implementation, analysis, visualization, and documentation. It demonstrates a complete workflow from target input to final report export. It uses open-source and modern technologies, presents a professional interface, and provides a foundation for future enhancements. Most importantly, it teaches that cybersecurity is a process of continuous visibility and prioritization. Networks change, vulnerabilities evolve, and defenders must keep mapping, measuring, and improving. This project is a practical step toward that mindset and a strong foundation for further development in network security engineering.
""")


def add_bibliography(doc: Document) -> None:
    doc.add_heading("12. BIBLIOGRAPHY", level=1)
    for item in BIBLIOGRAPHY:
        add_para(doc, item, after=4)


def main() -> None:
    doc = Document()
    setup_styles(doc)
    add_footer(doc.sections[0], "Network Topology Mapper & Vulnerability Analyzer")

    add_cover(doc)
    add_front_matter(doc)
    add_toc(doc)

    doc.add_heading("6. INTRODUCTION OF ORGANIZATION", level=1)
    add_body(doc, ORG_SECTION)

    doc.add_heading("7. INTERNSHIP CONTENT", level=1)
    add_body(doc, INTERNSHIP_CONTENT)

    add_project_details(doc)
    add_testing(doc)
    add_results(doc)
    add_conclusion(doc)
    add_bibliography(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    text_parts: list[str] = []
    for paragraph in doc.paragraphs:
        text_parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)
    print(OUT)
    print("word_count", words("\n".join(text_parts)))


if __name__ == "__main__":
    main()
